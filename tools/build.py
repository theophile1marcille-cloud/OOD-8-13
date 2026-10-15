"""Construction du Guide du marché CT à partir de content.yaml (rejouable).

    python3 tools/build.py [--content livrables/content.yaml] [--out livrables/Guide_du_marche_CT_2026.pptx]

- Part du .pptx de couverture (masque + thème « Guide du Marché CT Philips ») : toutes les
  slides utilisent ses mises en page ; la slide de couverture d'origine est conservée.
- Tout le texte vient de content.yaml ; ce script ne contient aucune donnée produit.
- Couleurs : uniquement la palette du thème (+ teintes calculées par mélange de deux couleurs du thème).
  meta.theme_overrides remplace des couleurs du thème (ex. accent1 = 0B5ED7) et est écrit dans le
  thème du fichier livré : la règle « palette du thème uniquement » reste vraie.
- Direction artistique « Verre » : alternance de slides sombres (ouvertures, chiffres clés, bandeaux
  versus) et claires à grille suisse (contenu), halos lumineux, cartes de verre dépoli (Liquid Glass) :
  dégradé blanc translucide + liseré + ombre douce, posé sur une copie floutée de la photo source.
- Les marqueurs [À COMPLÉTER] / [À VÉRIFIER] / [VISUEL À AJOUTER] sont stylés à part et
  recensés automatiquement dans A_COMPLETER.md ; les sources vont dans les notes.
"""
import argparse, copy, hashlib, re, sys
from pathlib import Path

import yaml
from PIL import Image, ImageEnhance, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
TODO, CHECK, VISUAL = "[À COMPLÉTER]", "[À VÉRIFIER]", "[VISUEL À AJOUTER]"
FONT = {False: "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        True: "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"}
SW, SH = 13.333, 7.5
MX = 0.6  # marge latérale


# --------------------------------------------------------------------------- palette
class Palette:
    def __init__(self, theme_xml):
        self.c = dict(re.findall(r"<a:(dk1|lt1|dk2|lt2|accent\d)><a:srgbClr val=\"(\w{6})\"", theme_xml))

    def __getitem__(self, key):
        """'accent1', 'dk1', ou mélange 'accent1/dk1/0.4' (40 % de la 2e couleur)."""
        if "/" in key:
            a, b, t = key.split("/")
            ca, cb, t = self.rgb(a), self.rgb(b), float(t)
            return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))
        return self.c[key]

    def rgb(self, key):
        h = self[key]
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))

    def lum(self, key):
        def ch(v):
            v /= 255
            return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
        r, g, b = (ch(v) for v in self.rgb(key))
        return 0.2126 * r + 0.7152 * g + 0.0722 * b

    def on(self, key):
        """Couleur de texte lisible sur un fond de la palette."""
        return "dk1" if self.lum(key) > 0.25 else "lt1"


# --------------------------------------------------------------------------- texte
_fonts = {}


def text_width(s, size, bold=False):
    k = (size, bold)
    if k not in _fonts:
        _fonts[k] = ImageFont.truetype(FONT[bold], int(size * 10))
    return _fonts[k].getlength(s) / 10 / 72  # pouces


def wrap_lines(text, size, width, bold=False):
    n = 0
    for para in str(text).split("\n"):
        line = ""
        for w in para.split(" "):
            t = (line + " " + w).strip()
            if text_width(t, size, bold) <= width or not line:
                line = t
            else:
                n += 1
                line = w
        n += 1
    return n


class Builder:
    def __init__(self, content_path, out_path):
        self.content_path = Path(content_path)
        self.content = yaml.safe_load(self.content_path.read_text())
        for i, sec in enumerate(self.content["sections"]):  # sections dans des fichiers séparés
            if "include" in sec:
                files = sec["include"] if isinstance(sec["include"], list) else [sec["include"]]
                parts = [yaml.safe_load((self.content_path.parent / f).read_text()) for f in files]
                merged = {**parts[0], "slides": [sl for part in parts for sl in part.get("slides", [])]}
                self.content["sections"][i] = {**merged, **{k: v for k, v in sec.items() if k != "include"}}
        meta = self.content["meta"]
        self.out_path = Path(out_path)
        self.img_root = ROOT / meta["images_root"]
        self.prs = Presentation(ROOT / meta["base_pptx"])
        theme = self.prs.slide_master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme").blob.decode()
        self.P = Palette(theme)
        self.P.c.update(meta.get("theme_overrides", {}))
        self.meta = meta
        self.dark = True
        self.cache = self.out_path.parent / "_cache"
        self.cache.mkdir(exist_ok=True)
        self.layout_blank = self.prs.slide_layouts.get_by_name("DEFAULT")
        self.layout_title = self.prs.slide_layouts.get_by_name("Title and body")
        self.ids, self.links, self.todos, self.warnings = {}, [], [], []
        self.slide_no = 0
        self._img_index = {}
        self.section = None

    # ------------------------------------------------------------------ primitives
    def color(self, key):
        return RGBColor.from_string(self.P[key])

    def rect(self, s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None,
             dash=False, line_w=1.0, name=None):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            sh.name = name
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = self.color(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = self.color(line)
            sh.line.width = Pt(line_w)
            if dash:
                sh.line.dash_style = 4  # dash
        else:
            sh.line.fill.background()
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = radius
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return sh

    def gradient(self, sh, stops, angle=0):
        """stops = [(pos 0-1, couleur, alpha 0-1)] ; dégradé linéaire natif."""
        spPr = sh._element.spPr
        for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
            for el in spPr.findall(qn(tag)):
                spPr.remove(el)
        g = etree.SubElement(spPr, qn("a:gradFill"), rotWithShape="1")
        lst = etree.SubElement(g, qn("a:gsLst"))
        for pos, col, alpha in stops:
            gs = etree.SubElement(lst, qn("a:gs"), pos=str(int(pos * 100000)))
            c = etree.SubElement(gs, qn("a:srgbClr"), val=self.P[col])
            if alpha < 1:
                etree.SubElement(c, qn("a:alpha"), val=str(int(alpha * 100000)))
        etree.SubElement(g, qn("a:lin"), ang=str(int(angle * 60000)), scaled="0")
        # gradFill doit précéder a:ln
        ln = spPr.find(qn("a:ln"))
        if ln is not None:
            spPr.remove(g)
            ln.addprevious(g)

    def text(self, s, x, y, w, h, runs, size=16, color="lt1", bold=False, align="l", anchor="t",
             italic=False, spacing=None, line=1.0, name=None, check=True, para_space=0):
        """runs : str, ou liste de paragraphes ; un paragraphe = str ou liste de (texte, {opts})."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        paras = runs if isinstance(runs, list) else [runs]
        plain = []
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = line
            if para_space and i:
                p.space_before = Pt(para_space)
            segs = para if isinstance(para, list) else [(para, {})]
            ptxt = ""
            for seg in segs:
                t, o = (seg, {}) if isinstance(seg, str) else seg
                for part in self.split_markers(str(t)):
                    r = p.add_run()
                    r.text = part
                    f = r.font
                    f.size = Pt(o.get("size", size))
                    f.bold = o.get("bold", bold)
                    f.italic = o.get("italic", italic)
                    col = o.get("color", color)
                    if part in (TODO, CHECK, VISUAL):
                        f.italic, f.bold = True, True
                        if part == CHECK:
                            col = "accent3" if self.dark else "accent1"
                        else:
                            col = "accent4" if self.dark else "accent6"
                        f.size = Pt(min(o.get("size", size), 14))
                    f.color.rgb = self.color(col)
                    if spacing:
                        r.font._element.set("spc", str(spacing))
                    ptxt += part
            plain.append((ptxt, max([size] + [o.get("size", size) for _, o in
                                              [(sg, {}) if isinstance(sg, str) else sg for sg in segs]]),
                          bold))
        if check:
            need = sum(wrap_lines(t, sz, w, b) * sz * 1.2 * line / 72 + (para_space / 72 if para_space else 0)
                       for t, sz, b in plain)
            if need > h + 0.05:
                self.warnings.append(f"slide {self.slide_no}: texte trop long ({need:.2f} > {h:.2f} in) : "
                                     f"{plain[0][0][:50]!r}")
        return tb

    @staticmethod
    def split_markers(t):
        return [p for p in re.split(r"(\[À COMPLÉTER\]|\[À VÉRIFIER\]|\[VISUEL À AJOUTER\])", t) if p]

    def image(self, s, ref, x, y, w, h, mode="cover", name=None, focus=(0.5, 0.5)):
        """cover : remplit la boîte en recadrant ; contain : tient dans la boîte. Jamais d'étirement."""
        path = self.resolve_image(ref)
        if path is None or not path.exists():
            self.warnings.append(f"slide {self.slide_no}: image absente {ref}")
            return self.visual_placeholder(s, x, y, w, h, ref)
        with Image.open(path) as im:
            iw, ih = im.size
        ir, br = iw / ih, w / h
        if mode == "contain":
            if ir > br:
                nw, nh = w, w / ir
            else:
                nw, nh = h * ir, h
            pic = s.shapes.add_picture(str(path), Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2),
                                       Inches(nw), Inches(nh))
            dpi = iw / nw
        else:
            pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
            if ir > br:  # trop large : recadrer gauche/droite
                keep = br / ir
                cut = 1 - keep
                pic.crop_left = cut * focus[0]
                pic.crop_right = cut * (1 - focus[0])
            else:
                keep = ir / br
                cut = 1 - keep
                pic.crop_top = cut * focus[1]
                pic.crop_bottom = cut * (1 - focus[1])
            dpi = (iw * (keep if ir > br else 1)) / w
        if dpi < 90:
            self.warnings.append(f"slide {self.slide_no}: résolution faible {ref} ({dpi:.0f} px/in)")
            self.todos.append((self.slide_no, self.cur_title,
                               f"Visuel {ref} : meilleure résolution souhaitable ({dpi:.0f} px/pouce à l'affichage ; "
                               f"c'est la meilleure version présente dans les sources)"))
        if name:
            pic.name = name
        pic.shadow.inherit = False
        return pic

    def resolve_image(self, ref):
        """« J-47 » = plus grande image de la page 47 du JFR ; « J-47#2 » = 2e plus grande ;
        sinon chemin relatif à analyse/."""
        m = re.fullmatch(r"([GJ])-(\d+)@(\d+)", str(ref))
        if m:  # n-ième image de la page en ordre de lecture (rangées de haut en bas, puis gauche à droite)
            imgs = self.page_images(m.group(1), int(m.group(2)))
            k = int(m.group(3)) - 1
            return imgs[k] if k < len(imgs) else None
        m = re.fullmatch(r"([GJ])-(\d+)(?:#(\d+))?", str(ref))
        if not m:
            return self.img_root / ref
        doc = {"G": "guide_2025", "J": "jfr_2026"}[m.group(1)]
        if doc not in self._img_index:
            import json
            self._img_index[doc] = json.loads((self.img_root / f"{doc}_images.json").read_text())
        page, rank = int(m.group(2)), int(m.group(3) or 1)
        cands = sorted((i for i in self._img_index[doc] if page in i["pages"]), key=lambda i: -i["w"] * i["h"])
        return self.img_root / cands[rank - 1]["file"] if len(cands) >= rank else None

    def page_images(self, kind, page, min_side=110):
        doc = {"G": "guide_2025", "J": "jfr_2026"}[kind]
        if doc not in self._img_index:
            import json
            self._img_index[doc] = json.loads((self.img_root / f"{doc}_images.json").read_text())
        c = [i for i in self._img_index[doc] if page in i["pages"] and min(i["w"], i["h"]) >= min_side]
        c.sort(key=lambda i: (round(i["bbox_pt"][1] / 90), i["bbox_pt"][0]))
        return [self.img_root / i["file"] for i in c]

    @staticmethod
    def round_pic(pic, radius):
        """Cadre arrondi natif (prstGeom roundRect) sur une image."""
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        for g in list(av):
            av.remove(g)
        etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(radius * 100000)}")

    def has_alpha(self, ref):
        p = self.resolve_image(ref)
        with Image.open(p) as im:
            return im.mode in ("RGBA", "LA", "P") and im.getextrema()[-1][0] < 250

    def visual_placeholder(self, s, x, y, w, h, hint=""):
        """Emplacement image réservé : verre + pointillés, libellé [VISUEL À AJOUTER]."""
        self.glass(s, x, y, w, h, radius=0.06)
        self.rect(s, x + 0.08, y + 0.08, w - 0.16, h - 0.16, line="accent4" if self.dark else "accent6",
                  dash=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, name="Visuel à ajouter")
        fg = "lt2" if self.dark else "dk2"
        if w < 1.6:
            self.text(s, x + 0.05, y, w - 0.1, h, [[("VISUEL", {"size": 9})], [("À AJOUTER", {"size": 9})]],
                      color="accent4" if self.dark else "accent6", bold=True, italic=True, align="c", anchor="m",
                      check=False)
        else:
            self.text(s, x + 0.1, y + h / 2 - 0.35, w - 0.2, 0.7, [[(VISUAL, {"size": 13})],
                                                                   [(hint, {"size": 10, "color": fg})]],
                      align="c", anchor="m", check=False)
        self.todos.append((self.slide_no, self.cur_title, VISUAL + (f" – {hint}" if hint else "")))

    def todo_box(self, s, x, y, w, h, label="", size=12):
        """Emplacement de donnée manquante : cadre pointillé gris + texte italique."""
        self.rect(s, x, y, w, h, line="accent4" if self.dark else "accent6", dash=True,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3, name="À compléter")
        self.text(s, x, y, w, h, TODO, size=size, align="c", anchor="m", check=False)

    def badge_new(self, s, x, y):
        b = self.rect(s, x, y, 1.0, 0.3, fill="accent1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5,
                      name="Badge Nouveau")
        self.glow(b, rad=9, alpha=0.55)
        self.text(s, x, y, 1.0, 0.3, "NOUVEAU", size=10, bold=True, color="lt1", align="c", anchor="m",
                  spacing=120, check=False)
        return b

    def token(self, s, x, y, d, label, size=11, filled=False):
        """Pictogramme homogène : pastille ronde + abréviation d'unité."""
        if filled:
            c = self.rect(s, x, y, d, d, fill="accent1", shape=MSO_SHAPE.OVAL, name="Picto")
            self.text(s, x, y, d, d, label, size=size, bold=True, color="lt1", align="c", anchor="m", check=False)
            return c
        col = "accent3" if self.dark else "accent1"
        c = self.rect(s, x, y, d, d, line=col, shape=MSO_SHAPE.OVAL, line_w=1.25, name="Picto")
        self.text(s, x, y, d, d, label, size=size, bold=True, color=col, align="c", anchor="m", check=False)
        return c

    # ------------------------------------------------------------------ effets « Verre »
    def _alpha(self, clr, a):
        if a < 1:
            etree.SubElement(clr, qn("a:alpha"), val=str(int(a * 100000)))

    def _spPr_fill(self, sh, el):
        spPr = sh._element.spPr
        for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
            for x in spPr.findall(qn(tag)):
                spPr.remove(x)
        geom = spPr.find(qn("a:prstGeom"))
        geom.addnext(el)

    def alpha_fill(self, sh, a):
        clr = sh._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        self._alpha(clr, a)

    def alpha_line(self, sh, a):
        clr = sh.line._get_or_add_ln().find(qn("a:solidFill")).find(qn("a:srgbClr"))
        self._alpha(clr, a)

    def _effects(self, sh):
        spPr = sh._element.spPr
        eff = spPr.find(qn("a:effectLst"))
        if eff is None:
            eff = etree.SubElement(spPr, qn("a:effectLst"))
        return eff

    def glow(self, sh, color="accent1", rad=16, alpha=0.5):
        g = etree.SubElement(self._effects(sh), qn("a:glow"), rad=str(int(rad * 12700)))
        self._alpha(etree.SubElement(g, qn("a:srgbClr"), val=self.P[color]), alpha)

    def soft_shadow(self, sh, alpha=0.3, blur=22, dist=6):
        eff = self._effects(sh)
        o = etree.Element(qn("a:outerShdw"), blurRad=str(int(blur * 12700)), dist=str(int(dist * 12700)),
                          dir="5400000", algn="t", rotWithShape="0")
        self._alpha(etree.SubElement(o, qn("a:srgbClr"), val=self.P["dk1"]), alpha)
        eff.insert(0, o)

    def halo(self, s, cx, cy, r, color="accent1", alpha=0.55):
        """Halo lumineux : disque à dégradé radial (couleur → transparent)."""
        sh = self.rect(s, cx - r, cy - r, 2 * r, 2 * r, fill=color, shape=MSO_SHAPE.OVAL, name="Halo")
        g = etree.Element(qn("a:gradFill"), rotWithShape="1")
        lst = etree.SubElement(g, qn("a:gsLst"))
        for pos, a in ((0, alpha), (45000, alpha * 0.45), (100000, 0)):
            gs = etree.SubElement(lst, qn("a:gs"), pos=str(pos))
            self._alpha(etree.SubElement(gs, qn("a:srgbClr"), val=self.P[color]), max(a, 0.0001))
        path = etree.SubElement(g, qn("a:path"), path="circle")
        etree.SubElement(path, qn("a:fillToRect"), l="50000", t="50000", r="50000", b="50000")
        self._spPr_fill(sh, g)
        return sh

    def glass(self, s, x, y, w, h, radius=0.1, under=None, strength=1.0, name="Verre"):
        """Carte de verre dépoli. under = (réf image, (bx, by, bw, bh), focus) : la photo visible sous la
        carte ; sa zone est recopiée floutée sous le verre (effet Liquid Glass)."""
        if under:
            self.frost(s, under, x, y, w, h, radius)
            if self.dark:
                veil = self.rect(s, x, y, w, h, fill="dk1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius,
                                 name="Verre – voile")
                self.alpha_fill(veil, getattr(self, "_veil", None) or 0.55)
        sh = self.rect(s, x, y, w, h, fill="lt1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius, name=name)
        top, bottom = ((0.24, 0.07) if self.dark else (0.88, 0.62))
        top, bottom = top * strength, bottom * strength
        g = etree.Element(qn("a:gradFill"), rotWithShape="1")
        lst = etree.SubElement(g, qn("a:gsLst"))
        for pos, a in ((0, top), (100000, bottom)):
            gs = etree.SubElement(lst, qn("a:gs"), pos=str(pos))
            self._alpha(etree.SubElement(gs, qn("a:srgbClr"), val=self.P["lt1"]), a)
        etree.SubElement(g, qn("a:lin"), ang=str(int(60 * 60000)), scaled="0")
        self._spPr_fill(sh, g)
        sh.line.color.rgb = self.color("lt1")
        sh.line.width = Pt(1)
        self.alpha_line(sh, 0.5 if self.dark else 0.95)
        self.soft_shadow(sh, alpha=0.35 if self.dark else 0.12)
        return sh

    def frost(self, s, under, x, y, w, h, radius):
        ref, (bx, by, bw, bh), focus = under
        path = self.resolve_image(ref)
        if path is None:
            return
        with Image.open(path) as im:
            im = im.convert("RGB")
            iw, ih = im.size
            sc = max(bw / iw, bh / ih)
            ox, oy = (iw * sc - bw) * focus[0], (ih * sc - bh) * focus[1]
            x0, y0 = max(x, bx), max(y, by)
            x1, y1 = min(x + w, bx + bw), min(y + h, by + bh)
            if x1 - x0 < 0.2 or y1 - y0 < 0.2:
                return
            box = tuple(int(v) for v in ((x0 - bx + ox) / sc, (y0 - by + oy) / sc,
                                         (x1 - bx + ox) / sc, (y1 - by + oy) / sc))
            key = hashlib.md5(f"{path}{box}".encode()).hexdigest()[:12]
            out = self.cache / f"frost_{key}.png"
            if not out.exists():
                pad = 40
                big = im.crop((max(box[0] - pad, 0), max(box[1] - pad, 0), min(box[2] + pad, iw),
                               min(box[3] + pad, ih)))
                big = big.filter(ImageFilter.GaussianBlur(max(iw, ih) / 70))
                big = ImageEnhance.Brightness(big).enhance(1.08)
                ox2, oy2 = box[0] - max(box[0] - pad, 0), box[1] - max(box[1] - pad, 0)
                big.crop((ox2, oy2, ox2 + box[2] - box[0], oy2 + box[3] - box[1])).save(out)
        pic = s.shapes.add_picture(str(out), Inches(x0), Inches(y0), Inches(x1 - x0), Inches(y1 - y0))
        pic.name = "Verre – flou de la photo source"
        if (x0, y0, x1, y1) == (x, y, x + w, y + h):
            self.round_pic(pic, radius)

    def frame(self, s, dark):
        """Fond de slide : sombre (bleu nuit) ou clair à grille suisse 12 colonnes."""
        self.dark = dark
        f = s.background.fill
        f.solid()
        f.fore_color.rgb = self.color("dk1" if dark else "lt1")
        if not dark:
            gw = (SW - 2 * MX) / 12
            for i in range(13):
                ln = self.rect(s, MX + i * gw, 0, 0.006, SH, fill="lt2", name="Grille")
                self.alpha_fill(ln, 0.5)

    def eyebrow(self, s, x, y, txt, w=6):
        """Petit sur-titre en capitales espacées."""
        self.text(s, x, y, w, 0.3, txt.upper(), size=10, bold=True, color="accent3" if self.dark else "accent1",
                  spacing=300, check=False)

    def link(self, shape, target):
        self.links.append((shape, target))

    # ------------------------------------------------------------------ cadre commun
    def new_slide(self, sd, layout="title"):
        s = self.prs.slides.add_slide(self.layout_title if layout == "title" else self.layout_blank)
        self.slide_no = len(self.prs.slides)
        self.cur_title = sd.get("title", "")
        for ph in list(s.placeholders):
            if ph.placeholder_format.type != 1:  # garder seulement le titre
                ph._element.getparent().remove(ph._element)
        if sd.get("id"):
            self.ids[sd["id"]] = s
        self.collect_markers(sd)
        self.notes(s, sd)
        return s

    def set_title(self, s, title, x=MX, y=0.78, w=SW - 2 * MX, h=0.75, size=30, color="lt1"):
        ph = s.shapes.title
        if ph is None:
            return self.text(s, x, y, w, h, title, size=size, bold=True, color=color)
        ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
        tf = ph.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = title
        r.font.size, r.font.bold = Pt(size), True
        r.font.color.rgb = self.color(color)
        lines = wrap_lines(title, size, w, True)
        self.title_lines = lines
        if lines * size * 1.15 / 72 > h + 0.05 and not getattr(self, "_allow_2l", False):
            self.warnings.append(f"slide {self.slide_no}: titre trop long {title!r}")
        return ph

    def chrome(self, s, sd, dark=None, show_section=True):
        """En-tête de section, pied de page, bouton sommaire, numéro, mention de confidentialité."""
        dark = self.dark if dark is None else dark
        if s.shapes.title is not None:  # le titre (placeholder du masque) passe au premier plan
            el = s.shapes.title._element
            tree = el.getparent()
            tree.remove(el)
            tree.append(el)
        sec = self.section
        fg, mute = ("lt1", "lt2") if dark else ("dk1", "dk2")
        if show_section and sec:
            label = f"{sec.get('number', 0):02d}  ·  {sec['short'].upper()}"
            self.rect(s, MX, 0.37, 0.15, 0.15, fill=sec["color"], shape=MSO_SHAPE.OVAL, name="Code section")
            t = self.text(s, MX + 0.25, 0.31, 6, 0.26, label, size=9, bold=True, color=mute, spacing=200,
                          check=False)
            self.link(t, sec["id"])
        if dark:
            b = self.glass(s, MX, 6.9, 1.3, 0.34, radius=0.5, strength=0.8, name="Retour sommaire")
        else:
            b = self.rect(s, MX, 6.9, 1.3, 0.34, line="dk1", line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                          radius=0.5, name="Retour sommaire")
        self.text(s, MX, 6.9, 1.3, 0.34, "‹  SOMMAIRE", size=9, bold=True, color=fg, align="c", anchor="m",
                  spacing=80, check=False)
        self.link(b, "sommaire")
        conf = self.meta["confidentiality"] + "  ·  " + self.meta["copyright"]
        for extra in sd.get("confidentiality_extra", []):
            conf = extra + "  ·  " + conf
        self.text(s, 2.15, 6.98, 8.6, 0.22, conf, size=8, color=mute, check=False)
        self.text(s, SW - MX - 2.2, 6.95, 1.0, 0.25, str(self.slide_no), size=9, color=mute, align="r",
                  check=False)
        self.text(s, SW - MX - 1.1, 6.93, 1.1, 0.3, "PHILIPS", size=13, bold=True, color=fg, align="r",
                  spacing=150, check=False)
        if sd.get("footnotes"):
            fn = sd["footnotes"]
            fn = fn if isinstance(fn, list) else [fn]
            fy = sd.get("footnotes_y", 6.2)
            self.text(s, MX, fy, sd.get("footnotes_w", SW - 2 * MX), 6.82 - fy, fn, size=8, color=mute,
                      anchor="b", line=1.0)

    def notes(self, s, sd):
        parts = []
        if sd.get("notes"):
            parts.append(sd["notes"].strip())
        src = sd.get("src")
        if src:
            src = src if isinstance(src, list) else [src]
            parts.append("Sources : " + " ; ".join(self.expand_src(x) for x in src))
        if parts:
            s.notes_slide.notes_text_frame.text = "\n\n".join(parts)

    def expand_src(self, ref):
        m = re.fullmatch(r"([GJC])-(\d+)(?:-(\d+))?", str(ref))
        if not m:
            return str(ref)
        doc = {"G": "guide_du_marche_2025 (PDF)", "J": "presentation_JFR_2026 (PDF)",
               "C": "couverture 2026 (.pptx)"}[m.group(1)]
        a, b = int(m.group(2)), m.group(3)
        rng = f"{a}" if not b else f"{a} à {b}"
        if m.group(1) == "G":
            sl = f"{a + 1}" if not b else f"{a + 1} à {int(b) + 1}"
            return f"{doc}, page {rng} (n° de slide affiché {sl})"
        return f"{doc}, page {rng}"

    def collect_markers(self, sd):
        if sd.get("type") == "versus":
            names = [sd["philips"]["name"], sd["competitor"]["name"]] + \
                    ([sd["secondary"]["name"]] if sd.get("secondary") else [])
            for row in sd["rows"]:
                for name, v in zip(names, row["values"]):
                    if str(v) in (TODO, CHECK):
                        self.todos.append((self.slide_no, sd.get("title", ""), f"{v} {row['label']} – {name}"))
            sd = {k: v for k, v in sd.items() if k != "rows"}
        if sd.get("type") == "table":
            for row in sd["rows"]:
                for head, v in zip(sd["columns"], row):
                    if str(v) in (TODO, CHECK):
                        self.todos.append((self.slide_no, sd.get("title", ""),
                                           f"{v} {' → '.join(str(c) for c in row if str(c) not in (TODO, CHECK))}"))
            sd = {k: v for k, v in sd.items() if k != "rows"}

        def walk(v, path):
            if isinstance(v, dict):
                for k, x in v.items():
                    if k not in ("notes", "src", "todo_notes"):
                        walk(x, path + [k])
            elif isinstance(v, list):
                for x in v:
                    walk(x, path)
            elif isinstance(v, str):
                if path and path[-1] in ("rows",) and isinstance(v, str) and v in (TODO, CHECK):
                    self.todos.append((self.slide_no, sd.get("title", ""), v))
                    return
                for m in (TODO, CHECK):
                    if m in v:
                        ctx = v.replace(m, "").strip(" :–-") or "/".join(str(p) for p in path)
                        self.todos.append((self.slide_no, sd.get("title", ""), f"{m} {ctx}"))
        walk(sd, [])
        for t in sd.get("todo_notes", []):
            self.todos.append((self.slide_no, sd.get("title", ""), t))

    # ------------------------------------------------------------------ types de slides
    def s_cover(self, sd):
        s = self.prs.slides[0]
        self.slide_no = 1
        self.cur_title = "Couverture"
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        for old, new in sd.get("replace", {}).items():
                            if old in r.text:
                                r.text = r.text.replace(old, new)
                        # le cyan d'origine (hors nouvelle charte) passe au bleu clair de la palette
                        if r.font.color and r.font.color.type is not None and str(r.font.color.rgb) == "7FD1FF":
                            r.font.color.rgb = self.color("accent3")
        self.notes(s, sd)
        self.ids["couverture"] = s

    def s_toc(self, sd):
        s = self.new_slide(sd)
        self.frame(s, dark=True)
        img_box = (7.4, 0, SW - 7.4, SH)
        self.image(s, sd["image"], *img_box, focus=(0.55, 0.5))
        ov = self.rect(s, 7.4, 0, 2.6, SH, fill="dk1")
        self.gradient(ov, [(0, "dk1", 1), (1, "dk1", 0)])
        self.halo(s, 2.2, 2.4, 3.2, alpha=0.6)
        self.halo(s, 6.6, 6.0, 2.6, alpha=0.4)
        self.eyebrow(s, MX, 0.55, "Guide du marché CT · Édition 2026")
        self.set_title(s, sd["title"], y=0.85, size=40, h=0.8)
        secs = [x for x in self.content["sections"] if not x.get("front")]
        cols, tw, th, gx, gy = 2, 3.15, 1.05, 0.22, 0.2
        for i, sec in enumerate(secs):
            cx, cy = MX + (i % cols) * (tw + gx), 1.95 + (i // cols) * (th + gy)
            under = ("" if cx + tw < 7.4 else None)
            card = self.glass(s, cx, cy, tw, th, radius=0.14, name=f"Sommaire {sec['title']}")
            n = self.rect(s, cx + 0.22, cy + 0.25, 0.55, 0.55, fill="accent1", shape=MSO_SHAPE.OVAL)
            self.glow(n, rad=10, alpha=0.6)
            self.text(s, cx + 0.22, cy + 0.25, 0.55, 0.55, f"{i + 1:02d}", size=13, bold=True, align="c",
                      anchor="m", check=False)
            self.text(s, cx + 0.95, cy + 0.18, tw - 1.05, 0.4, sec["title"], size=15, bold=True)
            self.text(s, cx + 0.95, cy + 0.56, tw - 1.05, 0.45, sec.get("tagline", ""), size=9.5, color="lt2")
            self.link(card, sec["id"])
        self.chrome(s, sd, show_section=False)

    def s_opener(self, sd):
        """Ouverture immersive (A + C) : panneau photo, numéro fantôme géant, halo, verre."""
        s = self.new_slide(sd)
        sec = self.section
        self.frame(s, dark=True)
        px = sd.get("photo_x", 6.9)
        img_box = (px, 0, SW - px, SH)
        focus = tuple(sd.get("focus", (0.6, 0.5)))
        self.image(s, sd["image"], *img_box, focus=focus)
        ov = self.rect(s, px, 0, 1.4, SH, fill="dk1")
        self.gradient(ov, [(0, "dk1", 1), (1, "dk1", 0)])
        self.halo(s, 2.6, 3.6, 3.6, alpha=0.55)
        self.deco_arcs(s, px, SH / 2, [3.2, 3.6, 4.1], start=95, end=265, alpha=0.5)
        num = sd.get("number")
        if num:
            self.text(s, MX - 0.15, 0.2, 7, 3.9, f"{num:02d}", size=260, bold=True, color="dk2", check=False)
        self.eyebrow(s, MX, 1.25, sd.get("eyebrow") or (f"Section {num:02d}" if num else sec["short"]))
        if sd.get("badge_new"):
            self.badge_new(s, MX, 1.7)
        tsize = sd.get("title_size", 66)
        while max(text_width(wd, tsize, True) for wd in sd["title"].split()) > 5.9 and tsize > 40:
            tsize -= 2
        self._allow_2l = True
        self.set_title(s, sd["title"], y=2.15, w=6.0, h=2.4, size=tsize)
        self._allow_2l = False
        ty = 2.15 + self.title_lines * tsize * 1.12 / 72 + 0.15
        if sd.get("subtitle"):
            self.text(s, MX, ty, 5.9, 0.9, sd["subtitle"], size=16, color="lt2")
        chips = sd.get("chapters", [])
        if chips:
            gw = px - MX - 0.35
            rows, cx_ = 1, 0.3
            for ch in chips:
                w_ = text_width(ch["label"], 12, True) + 0.5
                if cx_ + w_ > gw - 0.2:
                    rows, cx_ = rows + 1, 0.3
                cx_ += w_ + 0.14
            gh = 0.55 + 0.46 * rows
            gx, gy = MX, 6.55 - gh
            self.glass(s, gx, gy, gw, gh, radius=0.25 if rows == 1 else 0.15, name="Chapitres")
            self.text(s, gx + 0.3, gy + 0.12, 2, 0.25, "CHAPITRES", size=8, bold=True, color="lt2", spacing=250,
                      check=False)
            cx, cy = gx + 0.3, gy + 0.42
            for ch in chips:
                w = text_width(ch["label"], 12, True) + 0.5
                if cx + w > gx + gw - 0.2:
                    cx, cy = gx + 0.3, cy + 0.46
                c = self.rect(s, cx, cy, w, 0.36, fill="accent1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
                self.text(s, cx, cy, w, 0.36, ch["label"], size=12, bold=True, align="c", anchor="m", check=False)
                self.link(c, ch["link"])
                cx += w + 0.14
        self.chrome(s, sd, show_section=False)

    def s_xxl(self, sd):
        """Chiffre clé : un héros géant (A) ou une grille de cartes de verre (C)."""
        s = self.new_slide(sd)
        self.frame(s, dark=True)
        stats = sd["stats"]
        has_img = bool(sd.get("image"))
        area_w = 7.3 if has_img else SW - 2 * MX
        img_box = (8.1, 0, SW - 8.1, SH)
        focus = tuple(sd.get("focus", (0.5, 0.5)))
        if has_img:
            self.image(s, sd["image"], *img_box, focus=focus)
            ov = self.rect(s, 8.1, 0, 1.3, SH, fill="dk1")
            self.gradient(ov, [(0, "dk1", 1), (1, "dk1", 0)])
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.62, sd["eyebrow"])
        self._allow_2l = True
        self.set_title(s, sd["title"], y=0.95, w=area_w, h=1.3, size=30)
        self._allow_2l = False
        top = 0.95 + self.title_lines * 0.5 + 0.35
        if sd.get("lead"):
            self.text(s, MX, top - 0.15, area_w, 0.4, sd["lead"], size=14, color="lt2")
            top += 0.35
        if sd.get("hero", len(stats) == 1):
            st, rest = stats[0], stats[1:]
            self.halo(s, MX + 2.3, top + 1.4, 3.0, alpha=0.6)
            vs = sd.get("size", 150)
            self.text(s, MX - 0.1, top - 0.35, area_w, vs / 72 * 1.05,
                      [[(str(st["value"]), {"size": vs, "bold": True, "color": "lt1"}),
                        (f" {st.get('unit', '')}", {"size": int(vs * 0.32), "bold": True, "color": "accent3"})]],
                      check=False)
            y = top - 0.35 + vs / 72 * 1.0
            self.text(s, MX, y, area_w, 0.5, st["label"], size=22, bold=True)
            if st.get("detail"):
                self.text(s, MX, y + 0.5, area_w - 0.5, 0.7, st["detail"], size=14, color="lt2")
            if rest:
                cw = (area_w - 0.2 * (len(rest) - 1)) / len(rest)
                for i, r in enumerate(rest):
                    gx = MX + i * (cw + 0.2)
                    self.glass(s, gx, 5.3, cw, 0.95, radius=0.18)
                    self.text(s, gx + 0.2, 5.36, cw - 0.3, 0.45, [[(str(r["value"]), {"bold": True}),
                                                                  (f" {r.get('unit', '')}", {"size": 12})]],
                              size=22, check=False)
                    self.text(s, gx + 0.2, 5.8, cw - 0.3, 0.42, r["label"], size=10, color="lt2")
        else:
            n = len(stats)
            per_row = sd.get("per_row") or min(n, 4)
            rows = (n + per_row - 1) // per_row
            gap = 0.22
            cw = (area_w - gap * (per_row - 1)) / per_row
            avail = (6.15 if not sd.get("footnotes") else sd.get("footnotes_y", 6.2) - 0.15) - top
            ch = min(2.4, (avail - gap * (rows - 1)) / rows)
            self.halo(s, MX + area_w / 2, top + avail / 2, min(area_w, 6) / 2, alpha=0.45)
            vs = sd.get("size", 64 if ch > 1.9 else 48)
            for i, st in enumerate(stats):
                gx, gy = MX + (i % per_row) * (cw + gap), top + (i // per_row) * (ch + gap)
                under = (sd["image"], img_box, focus) if has_img and gx + cw > 8.1 else None
                self.glass(s, gx, gy, cw, ch, radius=0.12, under=under)
                val, unit = str(st["value"]), st.get("unit", "")
                v = vs
                while text_width(val, v, True) + text_width(unit, v * 0.4, True) + 0.15 > cw - 0.4 and v > 24:
                    v -= 2
                self.text(s, gx + 0.25, gy + 0.18, cw - 0.4, v / 72 * 1.15,
                          [[(val, {"size": v, "bold": True}), (f" {unit}" if unit else "",
                                                              {"size": int(v * 0.4), "bold": True,
                                                               "color": "accent3"})]], check=False)
                ly = gy + 0.2 + v / 72 * 1.15
                self.text(s, gx + 0.25, ly, cw - 0.45, gy + ch - ly - 0.1,
                          [[(st["label"], {"bold": True})]] + ([st["detail"]] if st.get("detail") else []),
                          size=12, color="lt1", para_space=3)
        if has_img:
            sd.setdefault("footnotes_w", 7.3)
        self.chrome(s, sd)

    def s_spec(self, sd):
        """Fiche technique (B) : photo cadrée + grille suisse de valeurs à filets fins."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.62, sd["eyebrow"])
        tsize = 30 if wrap_lines(sd["title"], 30, SW - 2 * MX - 1.2, True) == 1 else 22
        self.set_title(s, sd["title"], y=0.95, size=tsize, color="dk1", w=SW - 2 * MX - 1.2)
        if sd.get("badge_new"):
            self.badge_new(s, SW - MX - 1.0, 1.02)
        gw = (SW - 2 * MX) / 12
        ix, iy, iw, ih = MX, 1.9, 4 * gw - 0.15, 4.25
        if sd.get("image"):
            if self.has_alpha(sd["image"]):
                bg = self.rect(s, ix, iy, iw, ih, fill="dk1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
                self.halo(s, ix + iw / 2, iy + ih / 2, iw / 2, alpha=0.7)
                self.image(s, sd["image"], ix + 0.2, iy + 0.2, iw - 0.4, ih - 0.4, mode="contain")
            else:
                pic = self.image(s, sd["image"], ix, iy, iw, ih, focus=tuple(sd.get("focus", (0.5, 0.5))))
                self.round_pic(pic, 0.05)
        else:
            self.dark = True
            bg = self.rect(s, ix, iy, iw, ih, fill="dk1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
            self.visual_placeholder(s, ix + 0.3, iy + 0.3, iw - 0.6, ih - 0.6, sd.get("image_hint", ""))
            self.dark = False
        items = sd["items"]
        cols = sd.get("cols", 3)
        x0 = MX + 4 * gw + 0.15
        cw = (SW - MX - x0) / cols
        rows = (len(items) + cols - 1) // cols
        rh = min(1.12, 4.3 / rows)
        for i, it in enumerate(items):
            x, y = x0 + (i % cols) * cw, 1.9 + (i // cols) * rh
            self.rect(s, x, y, cw - 0.2, 0.012, fill="dk1")
            self.token(s, x, y + 0.17, 0.42, it.get("icon", "•"), size=8 if len(it.get("icon", "")) > 2 else 10)
            v = str(it["value"])
            label = it["label"]
            if CHECK in v:
                v = v.replace(CHECK, "").strip()
                label = f"{label}  {CHECK}"
            if v == TODO:
                self.todo_box(s, x + 0.55, y + 0.17, cw - 0.85, 0.34, size=10)
            else:
                self.text(s, x + 0.55, y + 0.12, cw - 0.75, 0.45, v, size=it.get("size", 19), bold=True,
                          color="dk1")
            self.text(s, x + 0.55, y + 0.55, cw - 0.75, rh - 0.6, label, size=10.5, color="dk2")
        self.chrome(s, sd)

    def s_cards(self, sd):
        """Gamme en bento (A) : une tuile focale bleue + tuiles de verre, photos produits."""
        s = self.new_slide(sd)
        self.frame(s, dark=True)
        self.halo(s, 3.2, 4.3, 3.5, alpha=0.6)
        self.halo(s, 10.5, 2.2, 3.0, alpha=0.35)
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.62, sd["eyebrow"])
        self.set_title(s, sd["title"], y=0.95, size=32)
        groups = sd["groups"]
        X, Y, W, H = MX, 1.95, SW - 2 * MX, 4.65
        lw = W * 0.42
        boxes = [(X, Y, lw, H)]
        rest = len(groups) - 1
        rx, rw = X + lw + 0.22, W - lw - 0.22
        heights = [H * 0.6, H * 0.4 - 0.22] if rest == 2 else [H]
        yy = Y
        for hgt in heights[:rest]:
            boxes.append((rx, yy, rw, hgt))
            yy += hgt + 0.22
        for gi, (g, (x, y, w, h)) in enumerate(zip(groups, boxes)):
            focal = gi == 0
            if focal:
                t = self.rect(s, x, y, w, h, fill="accent1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
                self.glow(t, rad=24, alpha=0.4)
            else:
                self.glass(s, x, y, w, h, radius=0.06)
            self.text(s, x + 0.3, y + 0.22, w - 0.6, 0.42, g["label"], size=20, bold=True)
            if g.get("pillar"):
                self.text(s, x + 0.3, y + 0.64, w - 0.6, 0.3, g["pillar"], size=11, italic=True,
                          color="lt1" if focal else "lt2")
            items = g["items"]
            if focal:
                hero = items[0]
                if hero.get("image"):
                    pic = self.image(s, hero["image"], x + 0.3, y + 1.1, w - 0.6, 2.05,
                                     focus=tuple(hero.get("focus", (0.5, 0.45))))
                    self.round_pic(pic, 0.06)
                ly = y + 3.3
                for it in items:
                    t = self.text(s, x + 0.3, ly, w - 1.8, 0.36, it["name"], size=15, bold=True)
                    if it.get("new"):
                        b = self.rect(s, x + w - 1.35, ly + 0.02, 1.05, 0.3, fill="lt1",
                                      shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
                        self.text(s, x + w - 1.35, ly + 0.02, 1.05, 0.3, "NOUVEAU", size=10, bold=True,
                                  color="accent1", align="c", anchor="m", spacing=120, check=False)
                    if it.get("link"):
                        self.link(t, it["link"])
                    ly += 0.4
            elif h > 2.2:
                n = len(items)
                cw = (w - 0.6 - 0.18 * (n - 1)) / n
                for i, it in enumerate(items):
                    cx = x + 0.3 + i * (cw + 0.18)
                    if it.get("image"):
                        pic = self.image(s, it["image"], cx, y + 1.1, cw, 1.15,
                                         focus=tuple(it.get("focus", (0.5, 0.55))))
                        self.round_pic(pic, 0.1)
                    else:
                        self.visual_placeholder(s, cx, y + 1.1, cw, 1.15, it.get("image_hint", ""))
                    t = self.text(s, cx, y + 2.35, cw, 0.3, it["name"], size=13, bold=True)
                    if it.get("sub"):
                        self.text(s, cx, y + 2.63, cw, 0.3, it["sub"], size=8, color="lt2")
                    if it.get("new"):
                        self.badge_new(s, cx + 0.06, y + 1.16)
                    if it.get("link"):
                        self.link(t, it["link"])
            else:
                colw = (w - 0.6) / 2
                for k, it in enumerate(items):
                    cx, cy = x + 0.3 + (k % 2) * colw, y + 0.9 + (k // 2) * 0.32
                    t = self.text(s, cx, cy, colw - 1.1, 0.3, it["name"], size=12, bold=True, check=False)
                    if it.get("new"):
                        self.rect(s, cx + colw - 1.05, cy + 0.03, 0.85, 0.24, fill="accent1",
                                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
                        self.text(s, cx + colw - 1.05, cy + 0.03, 0.85, 0.24, "NOUVEAU", size=8, bold=True,
                                  align="c", anchor="m", check=False)
                    if it.get("link"):
                        self.link(t, it["link"])
                if any(not it.get("image") for it in items):
                    self.todos.append((self.slide_no, sd["title"],
                                       f"{VISUAL} – {g['label']} : photos produits absentes des sources"))
        self.chrome(s, sd)

    def s_fullbleed(self, sd):
        """Grand visuel plein cadre + message sur carte de verre dépoli."""
        s = self.new_slide(sd)
        self.frame(s, dark=True)
        box = (0, 0, SW, SH)
        focus = tuple(sd.get("focus", (0.5, 0.5)))
        self.image(s, sd["image"], *box, focus=focus)
        side = sd.get("side", "l")
        ov = self.rect(s, 0, 0, SW, SH, fill="dk1")
        self.gradient(ov, [(0, "dk1", 0.55), (0.55, "dk1", 0.1), (1, "dk1", 0)], angle=0 if side == "l" else 180)
        x = MX if side == "l" else SW - MX - 6.0
        gy, gh = 1.15, 5.35
        self._veil = 0.72
        self.glass(s, x, gy, 6.0, gh, radius=0.07, under=(sd["image"], box, focus), strength=0.8)
        self._veil = None
        if sd.get("eyebrow"):
            self.eyebrow(s, x + 0.4, gy + 0.35, sd["eyebrow"])
        self._allow_2l = True
        self.set_title(s, sd["title"], x=x + 0.4, y=gy + 0.7, w=5.2, h=1.5, size=32)
        self._allow_2l = False
        y = gy + 0.75 + self.title_lines * 0.52 + 0.2
        if sd.get("big"):
            self.text(s, x + 0.35, y - 0.1, 5.3, 1.45,
                      [[(sd["big"]["value"], {"size": 96, "bold": True}),
                        (" " + sd["big"].get("unit", ""), {"size": 34, "bold": True, "color": "accent3"})]],
                      check=False)
            y += 1.45
        if sd.get("body"):
            self.text(s, x + 0.4, y, 5.2, gy + gh - y - 0.2, sd["body"], size=15)
        self.chrome(s, sd)

    def s_overview(self, sd):
        """Vue d'ensemble (B) : segments × constructeurs sur fond clair, Philips en bleu."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.5, sd["eyebrow"])
        self.set_title(s, sd["title"], y=0.78, h=0.6, size=26, color="dk1")
        vendors, segs = sd["vendors"], sd["segments"]
        x0, y0 = MX + 1.5, 1.5
        cw = (SW - MX - x0) / len(vendors)
        hdr = 0.42
        for j, v in enumerate(vendors):
            ph = v.get("philips")
            b = self.rect(s, x0 + j * cw + 0.04, y0, cw - 0.08, hdr, fill="accent1" if ph else "dk1",
                          shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
            if ph:
                self.glow(b, rad=8, alpha=0.45)
            self.text(s, x0 + j * cw, y0, cw, hdr, v["name"], size=12, bold=True, color="lt1", align="c",
                      anchor="m", check=False)
        avail = 6.75 - (y0 + hdr + 0.1)
        weights = [seg.get("weight", 1) for seg in segs]
        y = y0 + hdr + 0.1
        for seg, wt in zip(segs, weights):
            rh = avail * wt / sum(weights)
            self.rect(s, MX, y, SW - 2 * MX, 0.014, fill="dk1")
            self.text(s, MX, y + 0.08, 1.4, 0.32, seg["name"], size=14, bold=True, color="accent1")
            self.text(s, MX, y + 0.42, 1.38, rh - 0.5, seg.get("desc", ""), size=8, color="dk2", line=0.95)
            for j, v in enumerate(vendors):
                items = seg["cells"].get(v["key"], [])
                if not items:
                    continue
                n = len(items)
                ih = min(0.6, (rh - 0.16) / n - 0.05)
                for k, it in enumerate(items):
                    ix, iy = x0 + j * cw + 0.06, y + 0.08 + k * (ih + 0.05)
                    ph, new = v.get("philips"), it.get("new2026")
                    fill = "accent1" if new else ("accent3" if ph else "lt1")
                    box = self.rect(s, ix, iy, cw - 0.12, ih, fill=fill, line=None if (ph or new) else "lt2",
                                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15, name=it["name"])
                    if new:
                        self.glow(box, rad=8, alpha=0.4)
                    tcol = "lt1" if new else "dk1"
                    lines = [[(it["name"], {"bold": True, "size": 9.5})]]
                    if it.get("spec") and ih > 0.38:
                        lines.append([(it["spec"], {"size": 7.5, "color": "lt1" if new else "dk2"})])
                    self.text(s, ix + 0.08, iy, cw - 0.3, ih, lines, color=tcol, anchor="m", line=0.9, check=False)
                    if new:
                        self.text(s, ix + cw - 0.85, iy + 0.03, 0.66, 0.2, "NOUVEAU", size=6.5, bold=True,
                                  color="lt1", align="r", check=False)
                    if it.get("link"):
                        self.link(box, it["link"])
            y += rh
        self.chrome(s, sd)

    def _product_frame(self, s, x, y, w, h, c, under):
        self.glass(s, x, y, w, h, radius=0.1, under=under)
        if c.get("image"):
            if self.has_alpha(c["image"]):
                self.halo(s, x + w / 2, y + h / 2, min(w, h) / 2, alpha=0.5)
                self.image(s, c["image"], x + 0.15, y + 0.12, w - 0.3, h - 0.24, mode="contain")
            else:
                pic = self.image(s, c["image"], x + 0.1, y + 0.1, w - 0.2, h - 0.2,
                                 focus=tuple(c.get("focus", (0.5, 0.5))))
                self.round_pic(pic, 0.08)
        else:
            self.visual_placeholder(s, x + 0.1, y + 0.1, w - 0.2, h - 0.2, f"photo {c['name']}")

    def s_versus(self, sd):
        """Face-à-face (A + B) : bandeau sombre avec photos produits sous verre et « VS » lumineux,
        tableau clair à filets ; Philips en bleu, concurrent en gris."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        ph, cp, sec2 = sd["philips"], sd["competitor"], sd.get("secondary")
        band_h = 2.45
        self.dark = True
        self.rect(s, 0, 0, SW, band_h, fill="dk1", name="Bandeau versus")
        self.halo(s, SW / 2, band_h / 2 + 0.1, 1.2, alpha=0.75)
        for k, rr in enumerate((0.8, 1.05, 1.35)):
            self.arc(s, SW / 2, 1.37, rr, 300, 60, width=1.2, color="accent3", alpha=0.5 - k * 0.12)
            self.arc(s, SW / 2, 1.37, rr, 120, 240, width=1.2, color="accent3", alpha=0.5 - k * 0.12)
        self.eyebrow(s, MX, 0.3, sd.get("eyebrow", "Compétition · face-à-face"))
        fw, fh = 2.0, 1.55
        self._product_frame(s, MX, 0.7, fw, fh, ph, None)
        self._product_frame(s, SW - MX - fw, 0.7, fw, fh, cp, None)
        self.text(s, MX + fw + 0.25, 0.85, 3.0, 1.2, [[(ph["name"], {"bold": True, "size": 22})],
                                                       [(ph.get("vendor", "Philips"), {"size": 11, "color": "lt2"})]],
                  check=False)
        self.text(s, SW - MX - fw - 3.25, 0.85, 3.0, 1.2,
                  [[(cp["name"], {"bold": True, "size": 22})], [(cp.get("vendor", ""), {"size": 11,
                                                                                          "color": "lt2"})]],
                  align="r", check=False)
        if ph.get("new"):
            self.badge_new(s, MX + fw + 0.25, 1.95)
        vs = self.rect(s, SW / 2 - 0.55, 0.82, 1.1, 1.1, fill="accent1", shape=MSO_SHAPE.OVAL, name="VS")
        self.glow(vs, rad=22, alpha=0.6)
        self.text(s, SW / 2 - 0.55, 0.82, 1.1, 1.1, "VS", size=24, bold=True, align="c", anchor="m", check=False)
        if s.shapes.title is not None:  # titre accessible (lecteurs d'écran, plan), hors champ visuel
            self.set_title(s, sd["title"], x=MX, y=0.0, w=0.1, h=0.1, size=1, color="dk1")
            self.title_lines = 1
        self.dark = False
        # tableau
        cols = [ph, cp] + ([sec2] if sec2 else [])
        lab_w = 3.1
        tx = MX
        cw = (SW - 2 * MX - lab_w) / len(cols)
        xc = [tx + lab_w + i * cw for i in range(len(cols))]
        rows = sd["rows"]
        y0 = band_h + 0.32
        hdr = 0.36
        rh = min(0.42, (6.45 - y0 - hdr) / len(rows))
        self.rect(s, xc[0], y0 - 0.06, cw, hdr + rh * len(rows) + 0.06, fill="accent3", name="Colonne Philips")
        for i, c in enumerate(cols):
            self.text(s, xc[i] + 0.2, y0, cw - 0.3, hdr,
                      c["name"] + (" (colonne secondaire)" if i == 2 else ""), size=10, bold=True,
                      color="accent1" if i == 0 else "dk2", anchor="m", check=False)
        rh = min(0.42, (6.45 - y0 - hdr) / len(rows))
        y = y0 + hdr
        for row in rows:
            self.rect(s, tx, y, SW - 2 * MX, 0.01, fill="dk1")
            self.token(s, tx, y + (rh - 0.28) / 2, 0.28, row.get("icon", "•"), size=6)
            self.text(s, tx + 0.38, y, lab_w - 0.45, rh, row["label"], size=11, color="dk1", anchor="m",
                      check=False)
            adv = row.get("advantage") == "philips"
            for i, v in enumerate(row["values"]):
                v = str(v)
                if v == TODO:
                    self.todo_box(s, xc[i] + 0.2, y + 0.06, min(cw - 0.4, 2.2), rh - 0.12, size=9)
                    continue
                self.text(s, xc[i] + 0.2, y, cw - (0.75 if (i == 0 and adv) else 0.3), rh, v,
                          size=13 if len(v) < 22 else 10, bold=(i == 0), color="dk1" if i else "accent1",
                          anchor="m", check=False)
                if i == 0 and adv:
                    d = self.rect(s, xc[0] + cw - 0.48, y + (rh - 0.28) / 2, 0.28, 0.28, fill="accent1",
                                  shape=MSO_SHAPE.OVAL, name="Avantage Philips")
                    self.text(s, xc[0] + cw - 0.48, y + (rh - 0.28) / 2, 0.28, 0.28, "✓", size=10, bold=True,
                              color="lt1", align="c", anchor="m", check=False)
            y += rh
        self.rect(s, tx, y, SW - 2 * MX, 0.01, fill="dk1")
        ly = 6.52
        d = self.rect(s, tx, ly + 0.04, 0.2, 0.2, fill="accent1", shape=MSO_SHAPE.OVAL)
        self.text(s, tx + 0.28, ly, 3.2, 0.28, "Avantage Philips (donnée sourcée)", size=9, color="dk2")
        self.rect(s, tx + 3.3, ly + 0.03, 0.5, 0.22, line="accent6", dash=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                  radius=0.4)
        self.text(s, tx + 3.9, ly, 4, 0.28, "Donnée absente des sources : à compléter", size=9, color="dk2")
        if sd.get("next"):
            n = self.rect(s, SW - MX - 2.5, ly - 0.04, 2.5, 0.34, fill="accent1", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                          radius=0.5)
            self.text(s, SW - MX - 2.5, ly - 0.04, 2.5, 0.34, "Arguments & questions  ›", size=10, bold=True,
                      color="lt1", align="c", anchor="m", check=False)
            self.link(n, sd["next"])
        self.chrome(s, sd, show_section=False)

    def s_arguments(self, sd):
        """Arguments (B) + encadré « Questions à poser au client » sur panneau sombre lumineux (C)."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.62, sd["eyebrow"])
        self._allow_2l = True
        self.set_title(s, sd["title"], y=0.95, w=7.6, h=1.0, size=26, color="dk1")
        self._allow_2l = False
        args = sd["arguments"]
        aw = 7.4
        top = 1.0 + self.title_lines * 0.45 + 0.35
        ah = min(1.2, (6.6 - top) / len(args) - 0.1)
        for i, a in enumerate(args):
            y = top + i * (ah + 0.1)
            self.rect(s, MX, y, aw, 0.012, fill="dk1")
            self.token(s, MX, y + 0.2, 0.62, a.get("icon", str(i + 1)), size=10, filled=True)
            self.text(s, MX + 0.85, y + 0.14, aw - 0.9, 0.35, a["title"], size=15, bold=True, color="accent1")
            self.text(s, MX + 0.85, y + 0.48, aw - 0.9, ah - 0.52, a["text"], size=11, color="dk1", line=1.0)
        qx, qw = MX + aw + 0.35, SW - 2 * MX - aw - 0.35
        self.dark = True
        p = self.rect(s, qx, 0.62, qw, 5.98, fill="dk1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
        self.soft_shadow(p, alpha=0.25)
        self.halo(s, qx + qw * 0.6, 2.3, min(qw * 0.36, 1.4), alpha=0.65)
        self.glass(s, qx + 0.25, 1.65, qw - 0.5, 4.7, radius=0.06, strength=0.8)
        self.eyebrow(s, qx + 0.3, 0.9, "À poser au client")
        self.text(s, qx + 0.3, 1.15, qw - 0.6, 0.4, "Questions clés", size=18, bold=True)
        qs = sd.get("questions") or [TODO]
        self.text(s, qx + 0.5, 1.85, qw - 1.0, 4.4, [[("?  ", {"bold": True, "color": "accent3"}), (q, {})]
                                                     for q in qs], size=12, para_space=9)
        self.dark = False
        if sd.get("back"):
            b = self.rect(s, MX + aw - 1.9, 0.6, 1.9, 0.32, line="accent1", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                          radius=0.5)
            self.text(s, MX + aw - 1.9, 0.6, 1.9, 0.32, "‹  Face-à-face", size=10, bold=True, color="accent1",
                      align="c", anchor="m", check=False)
            self.link(b, sd["back"])
        self.chrome(s, sd)

    def s_procons(self, sd):
        """Versus 2025 conservé : points forts du concurrent / points d'attention (données d'origine)."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.62, sd["eyebrow"])
        self.set_title(s, sd["title"], y=0.95, size=26, color="dk1")
        cols = [(sd.get("pros_label", "Points forts du concurrent"), sd.get("pros", []), "accent6", "+"),
                (sd.get("cons_label", "Points d'attention"), sd.get("cons", []), "accent1", "!")]
        cw = (SW - 2 * MX - 0.4) / 2
        for i, (head, items, col, sym) in enumerate(cols):
            x = MX + i * (cw + 0.4)
            self.rect(s, x, 1.85, cw, 0.014, fill="dk1")
            t = self.rect(s, x, 2.05, 0.5, 0.5, fill=col, shape=MSO_SHAPE.OVAL)
            self.text(s, x, 2.05, 0.5, 0.5, sym, size=18, bold=True, color="lt1", align="c", anchor="m",
                      check=False)
            self.text(s, x + 0.7, 2.1, cw - 0.8, 0.45, head, size=17, bold=True, color=col)
            self.text(s, x, 2.8, cw - 0.2, 3.7, [[("›  ", {"bold": True, "color": col}), (t_, {})] for t_ in items],
                      size=13, color="dk1", para_space=7)
        if sd.get("back"):
            b = self.rect(s, SW - MX - 1.9, 0.98, 1.9, 0.32, line="accent1", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                          radius=0.5)
            self.text(s, SW - MX - 1.9, 0.98, 1.9, 0.32, "‹  Face-à-face", size=10, bold=True, color="accent1",
                      align="c", anchor="m", check=False)
            self.link(b, sd["back"])
        self.chrome(s, sd)

    def s_table(self, sd):
        """Tableau (B) : en-tête bleu nuit, filets fins."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.62, sd["eyebrow"])
        self.set_title(s, sd["title"], y=0.95, size=30, color="dk1")
        if sd.get("lead"):
            self.text(s, MX, 1.65, SW - 2 * MX, 0.45, sd["lead"], size=14, color="dk2")
        cols, rows = sd["columns"], sd["rows"]
        widths = sd.get("widths") or [1] * len(cols)
        tw = SW - 2 * MX
        ws = [tw * w / sum(widths) for w in widths]
        y, rh = sd.get("top", 2.3), sd.get("row_h", 0.6)
        self.rect(s, MX, y, tw, rh, fill="dk1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
        x = MX
        for c, w in zip(cols, ws):
            self.text(s, x + 0.2, y, w - 0.3, rh, c, size=13, bold=True, color="lt1", anchor="m", check=False)
            x += w
        for i, r in enumerate(rows):
            y += rh
            x = MX
            for j, (c, w) in enumerate(zip(r, ws)):
                self.text(s, x + 0.2, y, w - 0.3, rh, str(c), size=sd.get("size", 14), bold=(j == 0),
                          color="accent1" if j == 0 else "dk1", anchor="m", check=False)
                x += w
            self.rect(s, MX, y + rh, tw, 0.012, fill="dk1")
        self.chrome(s, sd)

    # ------------------------------------------------------------------ arcs et formes circulaires
    def arc(self, s, cx, cy, r, start, end, width=2.0, color="accent1", alpha=1.0, glow=False, name="Arc"):
        """Arc de cercle (trait) de start à end degrés, sens horaire, 0° = 3 h (convention PowerPoint)."""
        sh = s.shapes.add_shape(MSO_SHAPE.ARC, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
        sh.name = name
        sh.adjustments[0] = (start % 360) * 0.6
        sh.adjustments[1] = (end % 360) * 0.6
        sh.fill.background()
        sh.line.color.rgb = self.color(color)
        sh.line.width = Pt(width)
        if alpha < 1:
            self.alpha_line(sh, alpha)
        sh.shadow.inherit = False
        if glow:
            self.glow(sh, color=color, rad=10, alpha=0.5)
        return sh

    def band_arc(self, s, cx, cy, r, start, sweep, thick=0.18, color="accent1", alpha=1.0, glow=True):
        """Arc plein (blockArc) : jauge circulaire."""
        sh = s.shapes.add_shape(MSO_SHAPE.BLOCK_ARC, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
        sh.name = "Jauge"
        sh.adjustments[0] = (start % 360) * 0.6
        sh.adjustments[1] = ((start + sweep) % 360) * 0.6
        sh.adjustments[2] = thick / (2 * r)
        sh.fill.solid()
        sh.fill.fore_color.rgb = self.color(color)
        if alpha < 1:
            self.alpha_fill(sh, alpha)
        sh.line.fill.background()
        sh.shadow.inherit = False
        if glow:
            self.glow(sh, color=color, rad=12, alpha=0.55)
        return sh

    def deco_arcs(self, s, cx, cy, radii, start=200, end=340, alpha=0.35, width=1.25):
        for i, r in enumerate(radii):
            self.arc(s, cx, cy, r, start, end, width=width, color="accent3" if self.dark else "accent1",
                     alpha=alpha * (1 - i * 0.18), name="Arc décoratif")

    def oval_pic(self, pic):
        pic._element.spPr.find(qn("a:prstGeom")).set("prst", "ellipse")

    def bullets(self, s, x, y, w, h, items, size=13, color=None, dot=None, para_space=6):
        color = color or ("lt1" if self.dark else "dk1")
        dot = dot or ("accent3" if self.dark else "accent1")
        return self.text(s, x, y, w, h, [[("●  ", {"size": max(size - 5, 7), "color": dot}), (it, {})]
                                         for it in items], size=size, color=color, para_space=para_space)

    def head(self, s, sd, dark, title_w=None, size=30, y=0.95):
        """Sur-titre facultatif + titre, renvoie l'ordonnée sous le titre."""
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.62, sd["eyebrow"])
        self._allow_2l = True
        self.set_title(s, sd["title"], y=y, w=title_w or (SW - 2 * MX), h=1.3, size=size,
                       color="lt1" if dark else "dk1")
        self._allow_2l = False
        return y + self.title_lines * size * 1.18 / 72 + 0.25

    # ------------------------------------------------------------------ nouveaux types
    def s_gauges(self, sd):
        """Jauges circulaires (pourcentages) : arcs pleins lumineux sur fond sombre."""
        s = self.new_slide(sd)
        self.frame(s, dark=True)
        top = self.head(s, sd, True)
        if sd.get("lead"):
            self.text(s, MX, top - 0.1, SW - 2 * MX, 0.4, sd["lead"], size=14, color="lt2")
            top += 0.35
        stats = sd["stats"]
        n = len(stats)
        cw = (SW - 2 * MX) / n
        r = min(1.25, cw / 2 - 0.35)
        cy = top + r + 0.25
        self.halo(s, SW / 2, cy, 3.5, alpha=0.4)
        for i, st in enumerate(stats):
            cx = MX + cw * i + cw / 2
            pct = st.get("pct")
            ring = self.rect(s, cx - r, cy - r, 2 * r, 2 * r, line="lt1", shape=MSO_SHAPE.OVAL, line_w=10)
            self.alpha_line(ring, 0.12)
            if pct is not None:
                self.band_arc(s, cx, cy, r + 0.07, 270, max(min(pct, 100), 1) * 3.6 - 0.01, thick=0.2)
            else:
                self.glass(s, cx - r + 0.2, cy - r + 0.2, 2 * r - 0.4, 2 * r - 0.4, radius=0.5)
            val, unit = str(st["value"]), st.get("unit", "")
            vs = 44
            while text_width(val, vs, True) + text_width(unit, vs * 0.45, True) > 2 * r - 0.35 and vs > 18:
                vs -= 2
            self.text(s, cx - r, cy - 0.45, 2 * r, 0.9, [[(val, {"size": vs, "bold": True}),
                                                          (f" {unit}" if unit else "", {"size": int(vs * 0.45),
                                                                                        "bold": True,
                                                                                        "color": "accent3"})]],
                      align="c", anchor="m", check=False)
            ly = cy + r + 0.3
            self.text(s, cx - cw / 2 + 0.15, ly, cw - 0.3, 0.6, st["label"], size=14, bold=True, align="c")
            if st.get("detail"):
                self.text(s, cx - cw / 2 + 0.15, ly + 0.55, cw - 0.3, 6.1 - ly - 0.55 - (0.0), st["detail"], size=11,
                          color="lt2", align="c")
        self.chrome(s, sd)

    def s_arcflow(self, sd):
        """Étapes ou frise disposées le long d'un grand arc de cercle."""
        s = self.new_slide(sd)
        self.frame(s, dark=sd.get("dark", True))
        dark = self.dark
        top = self.head(s, sd, dark, size=28)
        if sd.get("lead"):
            self.text(s, MX, top - 0.12, SW - 2 * MX, 0.4, sd["lead"], size=13, color="lt2" if dark else "dk2")
        steps = sd["steps"]
        n = len(steps)
        import math
        R, cx, cy = 6.4, SW / 2, 8.55
        self.halo(s, cx, 3.0, 3.2, alpha=0.35 if dark else 0.18)
        self.arc(s, cx, cy, R, 200, 340, width=3, color="accent1", glow=dark, name="Arc frise")
        self.arc(s, cx, cy, R + 0.35, 205, 335, width=1, color="accent3" if dark else "accent1", alpha=0.4)
        a0, a1 = 138, 42
        for i, st in enumerate(steps):
            th = math.radians(a0 + (a1 - a0) * (i / (n - 1) if n > 1 else 0.5))
            x, y = cx + R * math.cos(th), cy - R * math.sin(th)
            d = 0.62
            node = self.rect(s, x - d / 2, y - d / 2, d, d, fill="accent1", shape=MSO_SHAPE.OVAL, name="Étape")
            if dark:
                self.glow(node, rad=12, alpha=0.6)
            self.text(s, x - d / 2, y - d / 2, d, d, st.get("marker", f"{i + 1}"), size=12 if
                      len(st.get("marker", "1")) < 4 else 9, bold=True, color="lt1", align="c", anchor="m", check=False)
            bw = min(2.6, (SW - 2 * MX) / n - 0.1)
            by = y + 0.45
            x = min(max(x, MX + bw / 2), SW - MX - bw / 2)
            self.text(s, x - bw / 2, by, bw, 0.4, st["title"], size=14, bold=True, align="c",
                      color="lt1" if dark else "dk1")
            if st.get("text"):
                txt = st["text"] if isinstance(st["text"], list) else [st["text"]]
                self.text(s, x - bw / 2, by + 0.42, bw, max(0.6, 6.7 - by - 0.5), txt, size=10.5,
                          color="lt2" if dark else "dk2", align="c", para_space=3)
        self.chrome(s, sd)

    def s_hub(self, sd):
        """Visuel central cerclé + éléments en couronne (cartes de verre)."""
        s = self.new_slide(sd)
        self.frame(s, dark=True)
        top = self.head(s, sd, True, size=28)
        cx, cy, r = SW / 2, max(top + 1.75, 3.95), 1.45
        self.halo(s, cx, cy, 2.8, alpha=0.6)
        for k, rr in enumerate([r + 0.25, r + 0.55, r + 0.95]):
            self.arc(s, cx, cy, rr, 0, 359.9, width=1, color="accent3", alpha=0.35 - k * 0.08)
        if sd.get("image"):
            pic = self.image(s, sd["image"], cx - r, cy - r, 2 * r, 2 * r, focus=tuple(sd.get("focus", (0.5, 0.5))))
            self.oval_pic(pic)
        else:
            self.text(s, cx - r, cy - 0.5, 2 * r, 1.0, sd.get("center", ""), size=20, bold=True, align="c", anchor="m")
        ring = self.rect(s, cx - r, cy - r, 2 * r, 2 * r, line="accent1", shape=MSO_SHAPE.OVAL, line_w=2.5)
        self.glow(ring, rad=14, alpha=0.6)
        items = sd["items"]
        left, right = items[: (len(items) + 1) // 2], items[(len(items) + 1) // 2:]
        gw, gap = 3.55, 0.15
        for side, group in ((-1, left), (1, right)):
            n = len(group)
            gh = min(1.25, (6.55 - top - gap * (n - 1)) / max(n, 1))
            total = n * gh + (n - 1) * gap
            y0 = cy - total / 2
            for k, it in enumerate(group):
                gx = MX if side < 0 else SW - MX - gw
                gy = max(top, y0) + k * (gh + gap)
                self.glass(s, gx, gy, gw, gh, radius=0.12)
                tl = wrap_lines(it["title"], 13, gw - 0.4, True)
                self.text(s, gx + 0.22, gy + 0.1, gw - 0.4, 0.25 * tl + 0.1, it["title"], size=13, bold=True,
                          color="accent3")
                if it.get("text"):
                    ty_ = gy + 0.14 + 0.24 * tl
                    self.text(s, gx + 0.22, ty_, gw - 0.4, gy + gh - ty_ - 0.05, it["text"], size=10.5 if gh > 1.1
                              else 10)
        self.chrome(s, sd)

    def s_gallery(self, sd):
        """Galerie clinique : 1 ou 2 cas par slide, chacun avec 1 à 3 images légendées (sans recadrage)."""
        s = self.new_slide(sd)
        self.frame(s, dark=True)
        cases = sd["cases"]
        if sd.get("eyebrow"):
            self.eyebrow(s, MX, 0.62, sd["eyebrow"])
        if len(cases) == 1 or sd.get("title"):
            top = self.head(s, {"title": sd.get("title") or cases[0]["title"]}, True, size=24)
        else:
            top = 0.95
            if s.shapes.title is not None:  # titre accessible (plan, lecteurs d'écran), hors champ visuel
                self.set_title(s, "Galerie – " + (sd.get("eyebrow") or "images cliniques"), x=MX, y=0.0, w=0.1,
                               h=0.1, size=1, color="dk1")
        n = len(cases)
        cw = (SW - 2 * MX - 0.3 * (n - 1)) / n
        for ci, c in enumerate(cases):
            if c.get("page") and not c.get("images"):  # toutes les images de la page source, dans l'ordre
                kind, pg = c["page"].split("-")
                paths = self.page_images(kind, int(pg))
                caps = c.get("captions", [])
                c["images"] = [{"image": str(p.relative_to(self.img_root)), "caption": caps[k] if k < len(caps) else ""}
                               for k, p in enumerate(paths)] or [{"hint": f"visuels de la page {c['page']}"}]
            x0 = MX + ci * (cw + 0.3)
            y = top
            if n > 1 or sd.get("title"):
                self.text(s, x0, y, cw, 0.62, c["title"], size=15 if n > 1 else 17, bold=True, check=False)
                y += 0.62 if len(c["title"]) > 52 * cw / 6 else 0.4
            if c.get("text"):
                self.text(s, x0, y, cw, 0.5, c["text"], size=10.5, color="lt2")
                y += 0.5
            imgs = c["images"]
            m = len(imgs)
            iw = (cw - 0.15 * (m - 1)) / m
            extra = (0.42 if any(i.get("caption") for i in imgs) else 0.1) + (0.35 if sd.get("footnotes") else 0)
            extra += 0.32 * bool(c.get("legend")) + 0.25 * bool(c.get("params"))
            ih = 6.2 - y - extra
            if c.get("legend") or c.get("params"):
                ly = y + ih + 0.08
                lines = []
                if c.get("legend"):
                    lines.append([("De gauche à droite : ", {"bold": True, "color": "accent3"}), (c["legend"], {})])
                if c.get("params"):
                    lines.append([(c["params"], {"color": "lt2", "size": 8.5})])
                self.text(s, x0, ly, cw, 0.6, lines, size=9.5, color="lt1", check=False)
            # taille réelle de chaque image ajustée à sa case, rangée centrée verticalement
            dims = []
            for im in imgs:
                pth = self.resolve_image(im["image"]) if im.get("image") else None
                if pth is not None and pth.exists():
                    with Image.open(pth) as pim:
                        r_ = pim.width / pim.height
                    w_, h_ = (iw, iw / r_) if r_ > iw / ih else (ih * r_, ih)
                else:
                    w_, h_ = iw, ih
                dims.append((w_, h_))
            rowh = max(h for _, h in dims)
            yy = y + (ih - rowh) / 2
            for k, (im, (w_, h_)) in enumerate(zip(imgs, dims)):
                ix = x0 + k * (iw + 0.15) + (iw - w_) / 2
                iy = yy + (rowh - h_) / 2
                if im.get("image") and self.resolve_image(im["image"]) is not None:
                    fr = self.rect(s, ix - 0.04, iy - 0.04, w_ + 0.08, h_ + 0.08, fill="dk2", line="accent3",
                                   shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, line_w=0.5)
                    self.alpha_line(fr, 0.4)
                    self.image(s, im["image"], ix, iy, w_, h_, mode="contain")
                else:
                    self.visual_placeholder(s, ix, iy, w_, h_, im.get("hint", ""))
                if im.get("caption"):
                    self.text(s, ix, y + ih + 0.06, iw, 0.36, im["caption"], size=9.5, color="lt2", align="c",
                              check=False)
        self.chrome(s, sd)

    def s_options(self, sd):
        """Options cliniques d'un produit : niveaux / packs + matrice par configuration."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        top = self.head(s, sd, False, size=26)
        groups = sd["groups"]
        sym = {"inclus": ("✓", "accent1", True), "option": ("+", "accent1", False), "non": ("–", "lt2", False)}

        def matrix(x, y, mx_w, m):
            cols, rows = m["cols"], m["rows"]
            lw = 1.05
            cwm = (mx_w - lw) / len(cols)
            for j, c in enumerate(cols):
                self.text(s, x + lw + j * cwm, y, cwm, 0.3, c, size=9, bold=True, color="dk2", align="c", check=False)
            for i, row in enumerate(rows):
                ry = y + 0.32 + i * 0.36
                self.rect(s, x, ry, mx_w, 0.01, fill="lt2")
                self.text(s, x, ry, lw, 0.36, row[0], size=10, bold=True, color="dk1", anchor="m", check=False)
                for j, v in enumerate(row[1:]):
                    g, col, filled = sym.get(v, (v, "dk1", False))
                    ox = x + lw + j * cwm + cwm / 2 - 0.13
                    if v in sym and v != "non":
                        self.rect(s, ox, ry + 0.05, 0.26, 0.26, fill=col if filled else None,
                                  line=None if filled else col, shape=MSO_SHAPE.OVAL, line_w=1.25)
                    self.text(s, ox, ry + 0.05, 0.26, 0.26, g, size=10, bold=True,
                              color="lt1" if filled else col, align="c", anchor="m", check=False)
            return y + 0.32 + len(rows) * 0.36

        avail_h = 6.35 - top
        big = [g for g in groups if g.get("levels")]
        small = [g for g in groups if not g.get("levels")]
        y = top
        for g in big:
            levels = g["levels"]
            self.token(s, MX, y, 0.42, g.get("icon", g["name"][:2]), size=9, filled=True)
            self.text(s, MX + 0.55, y + 0.04, 6, 0.36, g["name"], size=16, bold=True, color="dk1")
            mw = 4.2 if g.get("matrix") else 0
            lw = (SW - 2 * MX - mw - 0.3) / len(levels)
            ly = y + 0.55
            h = sd.get("level_h", 1.9 if small else avail_h - 0.6)
            for k, lv in enumerate(levels):
                lx = MX + k * lw
                self.rect(s, lx, ly, lw - 0.2, 0.012, fill="dk1")
                self.text(s, lx, ly + 0.08, lw - 0.25, 0.3, lv["name"], size=12, bold=True, color="accent1")
                self.bullets(s, lx, ly + 0.42, lw - 0.25, h - 0.45, lv["items"], size=10.5, para_space=2)
            if g.get("matrix"):
                matrix(SW - MX - mw, ly, mw, g["matrix"])
            y = ly + h + 0.1
        if small:
            sw_ = (SW - 2 * MX - 0.3 * (len(small) - 1)) / len(small)
            for k, g in enumerate(small):
                x = MX + k * (sw_ + 0.3)
                self.token(s, x, y, 0.42, g.get("icon", g["name"][:2]), size=9, filled=True)
                self.text(s, x + 0.55, y + 0.04, sw_ - 0.6, 0.36, g["name"], size=15, bold=True, color="dk1")
                mw = min(2.4, sw_ * 0.42) if g.get("matrix") else 0
                bh = 6.35 - y - 0.55
                self.bullets(s, x, y + 0.55, sw_ - mw - 0.2, bh, g.get("items", []), size=10.5, para_space=2)
                if g.get("matrix"):
                    matrix(x + sw_ - mw, y + 0.5, mw, g["matrix"])
        legend_y = 6.5
        for k, (lab, key) in enumerate((("Inclus", "inclus"), ("En option", "option"))):
            g, col, filled = sym[key]
            lx = MX + k * 1.6
            self.rect(s, lx, legend_y + 0.03, 0.22, 0.22, fill=col if filled else None, line=None if filled else col,
                      shape=MSO_SHAPE.OVAL, line_w=1)
            self.text(s, lx, legend_y + 0.03, 0.22, 0.22, g, size=8, bold=True, color="lt1" if filled else col,
                      align="c", anchor="m", check=False)
            self.text(s, lx + 0.3, legend_y, 1.3, 0.28, lab, size=9, color="dk2", anchor="m", check=False)
        self.chrome(s, sd)

    def s_list(self, sd):
        """Liste compacte (base installée) : en-tête produit + tableau multi-colonnes à filets fins."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        top = self.head(s, sd, False, title_w=8.6, size=26)
        if sd.get("image"):
            self.halo(s, SW - MX - 1.1, 1.15, 1.0, alpha=0.25)
            self.image(s, sd["image"], SW - MX - 2.2, 0.45, 2.2, 1.3, mode="contain")
        if sd.get("count") is not None:
            self.text(s, SW - MX - 4.6, 0.55, 2.3, 1.0, [[(str(sd["count"]), {"size": 44, "bold": True,
                                                                               "color": "accent1"})],
                                                          [(sd.get("count_label", "installés"), {"size": 10,
                                                                                                "color": "dk2"})]],
                      align="r", check=False)
        cols = sd["columns"]
        rows = sd["rows"]
        ncol = sd.get("blocks", 1 if len(rows) <= 12 else (2 if len(rows) <= 34 else 3))
        per = (len(rows) + ncol - 1) // ncol
        bw = (SW - 2 * MX - 0.3 * (ncol - 1)) / ncol
        widths = sd.get("widths") or [1] * len(cols)
        rh = min(0.42, (6.4 - top - 0.35) / max(per, 1))
        fs = 11 if rh >= 0.3 else (9 if rh >= 0.22 else 8)
        for b in range(ncol):
            x0 = MX + b * (bw + 0.3)
            ws = [bw * w / sum(widths) for w in widths]
            x = x0
            for c, w in zip(cols, ws):
                self.text(s, x, top, w - 0.05, 0.3, c, size=9, bold=True, color="accent1", spacing=80, check=False)
                x += w
            self.rect(s, x0, top + 0.3, bw, 0.014, fill="dk1")
            for i, r in enumerate(rows[b * per:(b + 1) * per]):
                y = top + 0.33 + i * rh
                if i % 2:
                    band = self.rect(s, x0, y, bw, rh, fill="accent3")
                    self.alpha_fill(band, 0.45)
                x = x0
                for j, (c, w) in enumerate(zip(r, ws)):
                    self.text(s, x + 0.05, y, w - 0.1, rh, str(c), size=fs, bold=(j == 0), color="dk1",
                              anchor="m", check=False)
                    x += w
        self.chrome(s, sd)

    def s_sites(self, sd):
        """Fiches sites de visite (2 par slide)."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        top = self.head(s, sd, False, size=24)
        sites = sd["sites"]
        n = len(sites)
        cw = (SW - 2 * MX - 0.3 * (n - 1)) / n
        for k, st in enumerate(sites):
            x0 = MX + k * (cw + 0.3)
            y0 = top
            card = self.rect(s, x0, y0, cw, 6.6 - y0, fill="lt1", line="lt2", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                             radius=0.03)
            self.soft_shadow(card, alpha=0.12)
            hd = self.rect(s, x0, y0, cw, 0.62, fill="dk1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
            self.text(s, x0 + 0.2, y0, cw - 2.4, 0.62, [[(st["name"], {"bold": True, "size": 14})],
                                                         [(st.get("city", ""), {"size": 9, "color": "lt2"})]],
                      color="lt1", anchor="m", check=False)
            sysw = 2.1
            b = self.rect(s, x0 + cw - sysw - 0.15, y0 + 0.14, sysw, 0.34, fill="accent1",
                          shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            self.text(s, x0 + cw - sysw - 0.15, y0 + 0.14, sysw, 0.34, st.get("system", ""), size=10, bold=True,
                      color="lt1", align="c", anchor="m", check=False)
            y = y0 + 0.75
            img_w = 1.9 if st.get("image") else 0
            if st.get("image"):
                pic = self.image(s, st["image"], x0 + cw - img_w - 0.15, y, img_w, 1.35)
                self.round_pic(pic, 0.08)
            colw = (cw - img_w - 0.45) / 2
            fields = [("CT", st.get("ct", [])), ("ISP", st.get("isp", [])), ("Infos", st.get("info", [])),
                      ("Spécialités", st.get("specialites", []))]
            for i, (lab, vals) in enumerate(fields):
                fx = x0 + 0.15 + (i % 2) * (colw + 0.15)
                fy = y + (i // 2) * 0.95
                self.text(s, fx, fy, colw, 0.25, lab.upper(), size=8, bold=True, color="accent1", spacing=100,
                          check=False)
                vals = vals if isinstance(vals, list) else [vals]
                self.text(s, fx, fy + 0.24, colw, 0.68, [v for v in vals] or ["—"], size=8.5, color="dk1", line=0.95,
                          check=False)
            y += 1.95
            if st.get("comments"):
                self.text(s, x0 + 0.15, y - 0.05, cw - 0.3, 0.32, [[("Commentaires : ", {"bold": True}),
                                                                    (st["comments"], {})]], size=8.5, color="dk1",
                          check=False)
                y += 0.3
            sched = st.get("schedule")
            if sched:
                days = sched["days"]
                slots = sched.get("slots", ["Matin", "Après-midi"])
                gx, gw_ = x0 + 0.15, cw - 0.3
                dw = (gw_ - 0.9) / len(days)
                self.text(s, gx, y, 0.9, 0.22, sched.get("label", ""), size=7.5, bold=True, color="dk2", check=False)
                for d_i, d in enumerate(days):
                    self.text(s, gx + 0.9 + d_i * dw, y, dw, 0.22, d, size=7.5, bold=True, color="dk2", align="c",
                              check=False)
                for r_i, row in enumerate(sched["rows"]):
                    ry = y + 0.24 + r_i * 0.27
                    self.text(s, gx, ry, 0.9, 0.25, row["name"], size=7.5, bold=True, color="dk1", anchor="m",
                              check=False)
                    for d_i, cell in enumerate(row["cells"]):
                        c = self.rect(s, gx + 0.9 + d_i * dw + 0.02, ry + 0.02, dw - 0.04, 0.22,
                                      fill="accent3" if cell else "lt2", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                                      radius=0.3)
                        self.text(s, gx + 0.9 + d_i * dw, ry + 0.02, dw, 0.22, cell or "", size=6.5, color="dk1",
                                  align="c", anchor="m", check=False)
                y = y + 0.3 + len(sched["rows"]) * 0.27
            contact = st.get("contact", [])
            if contact:
                ch = 6.5 - y
                g = self.rect(s, x0 + 0.15, y, cw - 0.3, ch, fill="accent3", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                              radius=0.08)
                self.alpha_fill(g, 0.55)
                self.text(s, x0 + 0.3, y + 0.06, cw - 0.6, ch - 0.1,
                          [[("CONTACT", {"bold": True, "size": 8, "color": "accent1"})]] + contact, size=8,
                          color="dk1", line=0.95, check=False)
        self.chrome(s, sd)

    def s_study(self, sd):
        """Publication : visuel de l'étude + chiffre clé + conclusions."""
        s = self.new_slide(sd)
        self.frame(s, dark=False)
        top = self.head(s, sd, False, size=24)
        iw = 5.6
        if sd.get("image"):
            bg = self.rect(s, MX, top, iw, 6.3 - top, fill="dk1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
            self.dark = True
            self.halo(s, MX + iw / 2, top + (6.3 - top) / 2, 2.2, alpha=0.4)
            self.dark = False
            self.image(s, sd["image"], MX + 0.15, top + 0.15, iw - 0.3, 6.3 - top - 0.3, mode="contain")
        x = MX + iw + 0.4
        w = SW - MX - x
        y = top
        if sd.get("eyebrow2"):
            self.text(s, x, y, w, 0.3, sd["eyebrow2"].upper(), size=9, bold=True, color="accent1", spacing=150)
            y += 0.35
        if sd.get("stat"):
            st = sd["stat"]
            self.text(s, x, y - 0.1, w, 1.2, [[(str(st["value"]), {"size": 66, "bold": True, "color": "accent1"}),
                                               (" " + st.get("unit", ""), {"size": 26, "bold": True,
                                                                           "color": "accent1"})]], check=False)
            y += 1.1
            self.text(s, x, y, w, 0.75, st["label"], size=14, bold=True, color="dk1")
            y += 0.8
        ref_h = 0.2 + 0.16 * wrap_lines(sd.get("reference", ""), 8.5, w) if sd.get("reference") else 0
        if sd.get("points"):
            ph = min(0.66, (6.3 - ref_h - y) / len(sd["points"]))
            for p in sd["points"]:
                self.rect(s, x, y, w, 0.012, fill="lt2")
                self.text(s, x, y + 0.06, w, ph - 0.06, p, size=11 if ph > 0.5 else 10, color="dk1")
                y += ph
        if sd.get("reference"):
            self.text(s, x, 6.3 - ref_h, w, ref_h, sd["reference"], size=8.5, italic=True, color="dk2", anchor="b")
        self.chrome(s, sd)

    def s_feature(self, sd):
        """Visuel + titre + points (≤ 5) ; clair ou sombre, visuel à gauche ou à droite."""
        s = self.new_slide(sd)
        dark = sd.get("dark", False)
        self.frame(s, dark=dark)
        side = sd.get("image_side", "r")
        iw = sd.get("image_w", 5.4)
        has = bool(sd.get("image")) or sd.get("image_hint")
        tx = MX if side == "r" or not has else MX + iw + 0.45
        tw = (SW - 2 * MX - (iw + 0.45 if has else 0))
        top = self.head(s, {**sd, "title": sd["title"]}, dark, title_w=tw if has else None, size=sd.get("title_size", 28))
        if tx != MX:  # titre à droite du visuel
            ph = s.shapes.title
            ph.left = Inches(tx)
        ix = SW - MX - iw if side == "r" else MX
        if sd.get("image"):
            iy, ih = 0.62 if not dark else 0.0, 6.0 if not dark else SH
            if sd.get("footnotes") and not dark:
                ih = min(ih, sd.get("footnotes_y", 6.45) - iy - 0.12)
            if dark:
                box = (ix - (MX if side == "r" else 0), 0, iw + MX, SH)
                self.image(s, sd["image"], *box, mode=sd.get("image_mode", "cover"),
                           focus=tuple(sd.get("focus", (0.5, 0.5))))
                ov = self.rect(s, box[0] if side == "r" else box[0] + box[2] - 1.2, 0, 1.2, SH, fill="dk1")
                self.gradient(ov, [(0, "dk1", 1), (1, "dk1", 0)], angle=0 if side == "r" else 180)
            else:
                if sd.get("image_mode") == "contain" or self.has_alpha(sd["image"]):
                    bg = self.rect(s, ix, iy, iw, ih, fill="dk1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
                    self.dark = True
                    self.halo(s, ix + iw / 2, iy + ih / 2, min(iw, ih) / 2, alpha=0.5)
                    self.dark = False
                    self.image(s, sd["image"], ix + 0.15, iy + 0.15, iw - 0.3, ih - 0.3, mode="contain")
                else:
                    pic = self.image(s, sd["image"], ix, iy, iw, ih, focus=tuple(sd.get("focus", (0.5, 0.5))))
                    self.round_pic(pic, 0.04)
        elif sd.get("image_hint"):
            self.visual_placeholder(s, ix, 0.9, iw, 5.3, sd["image_hint"])
        y = top
        if sd.get("lead"):
            self.text(s, tx, y - 0.05, tw, 0.9, sd["lead"], size=15, bold=True, color="accent3" if dark else "accent1")
            y += 0.1 + wrap_lines(sd["lead"], 15, tw, True) * 0.27
        pts = sd.get("points", [])
        if pts and isinstance(pts[0], dict):
            ph = min(1.15, (6.3 - y) / len(pts) - 0.08)
            for p in pts:
                self.rect(s, tx, y, tw, 0.012, fill="lt2" if not dark else "accent3")
                self.text(s, tx, y + 0.08, tw, 0.32, p["title"], size=13.5, bold=True,
                          color="accent3" if dark else "accent1")
                if p.get("text"):
                    self.text(s, tx, y + 0.4, tw, ph - 0.42, p["text"], size=11, color="lt1" if dark else "dk1")
                y += ph + 0.08
        elif pts:
            self.bullets(s, tx, y, tw, 6.3 - y, pts, size=sd.get("size", 13))
        if sd.get("body"):
            self.text(s, tx, y, tw, 6.3 - y, sd["body"], size=sd.get("size", 13), color="lt1" if dark else "dk1",
                      para_space=6)
        self.chrome(s, sd)

    def s_grid(self, sd):
        """Grille de 2 à 6 cartes (picto + titre + texte) ; claire par défaut, verre en sombre."""
        s = self.new_slide(sd)
        dark = sd.get("dark", False)
        self.frame(s, dark=dark)
        top = self.head(s, sd, dark, size=28)
        if sd.get("lead"):
            self.text(s, MX, top - 0.1, SW - 2 * MX, 0.45, sd["lead"], size=13, color="lt2" if dark else "dk2")
            top += 0.45
        cards = sd["cards"]
        n = len(cards)
        cols = sd.get("cols") or (n if n <= 4 else 3)
        rows = (n + cols - 1) // cols
        gap = 0.25
        cw = (SW - 2 * MX - gap * (cols - 1)) / cols
        bottom = sd.get("bottom", 6.4)
        ch = min(sd.get("card_h", 2.9), (bottom - top - gap * (rows - 1)) / rows)
        if dark:
            self.halo(s, SW / 2, top + (bottom - top) / 2, 3.2, alpha=0.45)
        for i, c in enumerate(cards):
            x, y = MX + (i % cols) * (cw + gap), top + (i // cols) * (ch + gap)
            if dark:
                self.glass(s, x, y, cw, ch, radius=0.07)
            else:
                card = self.rect(s, x, y, cw, ch, fill="lt1", line="lt2", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                                 radius=0.05)
                self.soft_shadow(card, alpha=0.1)
            yy = y + 0.25
            if c.get("image"):
                pic = self.image(s, c["image"], x + 0.2, yy, cw - 0.4, min(1.6, ch * 0.42),
                                 mode=c.get("image_mode", "cover"))
                if c.get("image_mode") != "contain":
                    self.round_pic(pic, 0.06)
                yy += min(1.6, ch * 0.42) + 0.15
            elif c.get("icon"):
                self.token(s, x + 0.25, yy, 0.55, c["icon"], size=9 if len(c["icon"]) > 3 else 11, filled=True)
                yy += 0.7
            self.text(s, x + 0.25, yy, cw - 0.5, 0.6, c["title"], size=14, bold=True,
                      color="accent3" if dark else "accent1")
            yy += 0.25 + wrap_lines(c["title"], 14, cw - 0.5, True) * 0.25
            if c.get("text"):
                body = c["text"] if isinstance(c["text"], list) else [c["text"]]
                self.text(s, x + 0.25, yy, cw - 0.5, y + ch - yy - 0.15, body, size=c.get("size", 11),
                          color="lt1" if dark else "dk1", para_space=3)
        self.chrome(s, sd)

    # ------------------------------------------------------------------ assemblage
    def build(self):
        for i, sec in enumerate([x for x in self.content["sections"] if not x.get("front")], start=1):
            sec.setdefault("number", i)
        for sec in self.content["sections"]:
            sec.setdefault("short", sec["title"])
        for sec in self.content["sections"]:
            self.section = None if sec.get("front") else sec
            for j, sd in enumerate(sec["slides"]):
                if j == 0 and not sec.get("front"):
                    sd.setdefault("id", sec["id"])
                getattr(self, "s_" + sd["type"])(sd)
        for shape, target in self.links:
            if target not in self.ids:
                self.warnings.append(f"lien vers « {target} » sans cible")
                continue
            shape.click_action.target_slide = self.ids[target]
        self.write_theme()
        self.prs.core_properties.title = self.meta["title"]
        self.prs.save(self.out_path)
        self.write_todo()
        print(f"{len(self.prs.slides)} slides → {self.out_path}")
        for w in self.warnings:
            print("ATTENTION", w)
        return not self.warnings

    def write_theme(self):
        """Écrit les couleurs remplacées (meta.theme_overrides) dans le thème du fichier livré."""
        ov = self.meta.get("theme_overrides", {})
        if not ov:
            return
        part = self.prs.slide_master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
        xml = part.blob.decode()
        for k, v in ov.items():
            xml = re.sub(rf"(<a:{k}><a:srgbClr val=\")\w{{6}}", rf"\g<1>{v}", xml)
        part._blob = xml.encode()

    def write_todo(self):
        path = self.out_path.parent / "A_COMPLETER.md"
        lines = ["# À compléter – Guide du marché CT 2026", "",
                 "Généré automatiquement par `tools/build.py` à partir de `content.yaml` "
                 "(marqueurs [À COMPLÉTER], [À VÉRIFIER], [VISUEL À AJOUTER]).", ""]
        lines += self.content.get("todo_preamble", [])
        lines += ["", "| Slide | Titre | Élément |", "|---|---|---|"]
        seen = set()
        for n, t, item in self.todos:
            key = (n, item)
            if key in seen:
                continue
            seen.add(key)
            lines.append(f"| {n} | {t} | {item} |")
        path.write_text("\n".join(lines) + "\n")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--content", default=str(ROOT / "livrables" / "content.yaml"))
    ap.add_argument("--out", default=str(ROOT / "livrables" / "Guide_du_marche_CT_2026.pptx"))
    a = ap.parse_args()
    ok = Builder(a.content, a.out).build()
    sys.exit(0 if ok else 1)
