"""Prévisualisations de direction artistique (méthode « show, don't tell » de frontend-slides).

Trois directions, appliquées aux mêmes 4 slides représentatives, dans la palette Philips du thème :
  A « Signal »   – Bold Signal / bold-editorial : bleu nuit plein, typographie XXL, une carte focale cyan.
  B « Suisse »   – Swiss Modern / corporate : fond clair, grille visible, blocs bleus, filets fins.
  C « Anneau »   – Neon Cyber / dark-atmospheric : motif d'anneaux concentriques (statif du scanner, comme
                   la couverture), halos lumineux, photos cerclées.
Mises en page issues de baoyu-slide-deck : title-hero, bento-grid, key-stat, binary-comparison /
comparison-matrix. Densité : limites frontend-slides (≤ 6 cartes par grille, 1 idée par slide).

    python3 tools/previews.py   →  livrables/variantes_DA.pptx
"""
import sys
from pathlib import Path

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).parent))
from build import Builder, MX, SH, SW, text_width  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent

GAMME = [  # J-4 ; images : page de chaque produit
    ("Scanners spectraux", "Une innovation fiable et accessible",
     [("Verida", True, "J-44#2"), ("Spectral CT 7500", False, "G-5"), ("IQon CE", False, "G-16")]),
    ("Scanners conventionnels", "Flux de travail intelligents optimisés par l'IA",
     [("Rembra", True, "J-47"), ("CT 5300", False, "J-38"), ("CT 5100", False, "J-35"), ("CT 3500", False, "J-25")]),
    ("Radiothérapie", "Favoriser des soins en toute confiance",
     [("Rembra RT", True, None), ("Areta RT", True, None), ("Spectral CT 7500 RT", False, None),
      ("Verida RT", True, None)]),
]
VS_ROWS = [  # G-216 (versus 2025 conservé)
    ("Générateur", "120 kW", "108 kW", True),
    ("Couverture du détecteur", "8 cm", "16 cm", False),
    ("Temps de rotation", "0,27 s", "0,23 s", False),
    ("Vitesse de table", "600 mm/s", "437 mm/s", True),
    ("Gamme spectrale", "40 – 200 keV", "40 – 140 keV", True),
    ("Capacité calorifique", "30 MHU", "8 MHU", True),
]


class Previews(Builder):
    def __init__(self):
        content = {"meta": {"title": "Variantes DA", "base_pptx": "sources/guide_du_marche_2025.pptx",
                            "images_root": "analyse", "confidentiality": "Philips internal use only",
                            "copyright": "© 2026 Koninklijke Philips N.V. All rights reserved."},
                   "sections": []}
        self.content = content
        self.meta = content["meta"]
        self.out_path = ROOT / "livrables" / "variantes_DA.pptx"
        from pptx import Presentation
        from build import Palette
        self.img_root = ROOT / "analyse"
        self.prs = Presentation(ROOT / self.meta["base_pptx"])
        theme = self.prs.slide_master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme").blob.decode()
        self.P = Palette(theme)
        self.layout_blank = self.prs.slide_layouts.get_by_name("DEFAULT")
        self.layout_title = self.prs.slide_layouts.get_by_name("Title and body")
        self.ids, self.links, self.todos, self.warnings = {}, [], [], []
        self.slide_no, self._img_index, self.cur_title = 0, {}, ""
        self.section = {"id": "gamme", "short": "Gamme Philips", "color": "accent1"}

    # ------------------------------------------------------------ effets (XML natif)
    def alpha_line(self, sh, alpha):
        ln = sh.line._get_or_add_ln()
        fill = ln.find(qn("a:solidFill"))
        clr = fill.find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))

    def alpha_fill(self, sh, alpha):
        clr = sh._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))

    def glow(self, sh, color="accent1", rad=18, alpha=0.45):
        spPr = sh._element.spPr
        eff = spPr.find(qn("a:effectLst"))
        if eff is None:
            eff = etree.SubElement(spPr, qn("a:effectLst"))
        g = etree.SubElement(eff, qn("a:glow"), rad=str(int(rad * 12700)))
        c = etree.SubElement(g, qn("a:srgbClr"), val=self.P[color])
        etree.SubElement(c, qn("a:alpha"), val=str(int(alpha * 100000)))

    def text_glow(self, tb, color="accent1", rad=10, alpha=0.6):
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                rPr = r._r.get_or_add_rPr()
                eff = etree.SubElement(rPr, qn("a:effectLst"))
                g = etree.SubElement(eff, qn("a:glow"), rad=str(int(rad * 12700)))
                c = etree.SubElement(g, qn("a:srgbClr"), val=self.P[color])
                etree.SubElement(c, qn("a:alpha"), val=str(int(alpha * 100000)))
                # effectLst doit suivre solidFill et précéder latin/ea/cs
                rPr.remove(eff)
                fill = rPr.find(qn("a:solidFill"))
                (fill.addnext(eff) if fill is not None else rPr.insert(0, eff))

    def rings(self, s, cx, cy, radii, color="accent1", alpha=0.25, width=1.0):
        for r in radii:
            o = self.rect(s, cx - r, cy - r, 2 * r, 2 * r, line=color, shape=MSO_SHAPE.OVAL, line_w=width)
            self.alpha_line(o, alpha)

    def oval_pic(self, pic):
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "ellipse")

    def grid(self, s, color="lt2", cols=12, alpha=0.35):
        gw = (SW - 2 * MX) / cols
        for i in range(cols + 1):
            ln = self.rect(s, MX + i * gw, 0, 0.005, SH, fill=color)
            self.alpha_fill(ln, alpha)

    def blank(self, title=""):
        s = self.prs.slides.add_slide(self.layout_title)
        self.slide_no = len(self.prs.slides)
        self.cur_title = title
        for ph in list(s.placeholders):
            if ph.placeholder_format.type != 1:
                ph._element.getparent().remove(ph._element)
        return s

    def bg(self, s, color):
        f = s.background.fill
        f.solid()
        from pptx.dml.color import RGBColor
        f.fore_color.rgb = RGBColor.from_string(self.P[color])

    def label(self, s, x, y, txt, color="accent1"):
        self.text(s, x, y, 6, 0.3, txt.upper(), size=10, bold=True, color=color, spacing=300, check=False)

    def variant_tag(self, s, txt, dark=True):
        self.text(s, SW - MX - 4.5, 0.3, 4.5, 0.3, txt, size=9, bold=True, color="lt2" if dark else "accent6",
                  align="r", check=False)

    def nouveau(self, s, x, y, light=False):
        b = self.badge_new(s, x, y)
        return b

    # ============================================================ A « Signal »
    def a_opener(self):
        s = self.blank("A1")
        self.bg(s, "dk1")
        self.image(s, "J-44", 8.35, 0, SW - 8.35, SH, focus=(0.55, 0.5))
        self.text(s, MX - 0.1, 0.55, 8, 3.6, "01", size=250, bold=True, color="dk2", check=False)
        self.label(s, MX, 0.62, "Section 01")
        self.set_title(s, "Gamme\nPhilips", y=2.1, w=7.5, h=2.9, size=80)
        self.text(s, MX, 5.05, 6.6, 0.8, "Des solutions innovantes, simples à utiliser et qui font une réelle "
                  "différence", size=16, color="lt2")
        card = self.rect(s, MX, 5.85, 7.3, 0.8, fill="accent1", shape=MSO_SHAPE.RECTANGLE)
        self.text(s, MX + 0.3, 5.85, 6.9, 0.8, [[("Verida", {"bold": True}), ("   ·   Rembra   ·   CT 5300   ·   "
                                                                        "CT 5100   ·   CT 3500   ›", {})]],
                  size=14, color="dk1", anchor="m", check=False)
        self.variant_tag(s, "A · SIGNAL")
        self.chrome(s, {}, show_section=False)

    def a_bento(self):
        s = self.blank("A2")
        self.bg(s, "dk1")
        self.label(s, MX, 0.55, "Gamme Philips")
        self.set_title(s, "Trois familles, une même intelligence", y=0.9, size=34, h=0.7)
        # bento : grande tuile spectrale (focale cyan), tuile conventionnelle, tuile RT
        tiles = [(MX, 1.95, 5.6, 4.6, "accent1", GAMME[0]), (MX + 5.85, 1.95, 6.2, 2.85, "dk2", GAMME[1]),
                 (MX + 5.85, 5.05, 6.2, 1.5, "dk2", GAMME[2])]
        for x, y, w, h, fill, (fam, pillar, items) in tiles:
            self.rect(s, x, y, w, h, fill=fill)
            fg = "dk1" if fill == "accent1" else "lt1"
            self.text(s, x + 0.3, y + 0.25, w - 0.6, 0.45, fam, size=20, bold=True, color=fg)
            if h > 2:
                self.text(s, x + 0.3, y + 0.72, w - 0.6, 0.35, pillar, size=11, color=fg, italic=True)
            if fill == "accent1":
                self.image(s, "J-44#2", x + 0.3, y + 1.25, w - 0.6, 2.2, focus=(0.5, 0.45))
                yy = y + 3.6
                for name, new, _ in items:
                    self.text(s, x + 0.3, yy, 3, 0.3, name, size=14, bold=True, color="dk1")
                    if new:
                        self.rect(s, x + 2.0, yy + 0.02, 1.0, 0.27, fill="dk1")
                        self.text(s, x + 2.0, yy + 0.02, 1.0, 0.27, "NOUVEAU", size=9, bold=True, color="accent1",
                                  align="c", anchor="m", spacing=100, check=False)
                    yy += 0.32
            elif h > 2:
                n = len(items)
                cw = (w - 0.6 - 0.15 * (n - 1)) / n
                for i, (name, new, img) in enumerate(items):
                    cx = x + 0.3 + i * (cw + 0.15)
                    self.image(s, img, cx, y + 1.2, cw, 1.05, focus=(0.5, 0.55))
                    self.text(s, cx, y + 2.33, cw, 0.3, name, size=12, bold=True)
                    if new:
                        self.badge_new(s, cx + 0.06, y + 1.27)
            else:
                names = "   ·   ".join(n + (" ●" if new else "") for n, new, _ in items)
                self.text(s, x + 0.3, y + 0.8, w - 0.6, 0.4, names, size=12, color="lt2")
        self.variant_tag(s, "A · SIGNAL")
        self.chrome(s, {}, show_section=False)

    def a_stat(self):
        s = self.blank("A3")
        self.bg(s, "dk1")
        self.image(s, "J-50", 8.8, 0, SW - 8.8, SH, focus=(0.5, 0.5))
        self.label(s, MX, 0.55, "Rembra · Nouveau")
        self.text(s, MX - 0.1, 0.95, 8.2, 3.6, [[("85", {"size": 240, "bold": True, "color": "lt1"}),
                                                  (" cm", {"size": 72, "bold": True, "color": "accent1"})]],
                  check=False)
        self.set_title(s, "Le tunnel le plus large de sa catégorie", y=4.45, w=7.5, h=0.6, size=26)
        self.text(s, MX, 5.15, 7.2, 0.8, "Confort, accessibilité et flexibilité d'utilisation pour le patient "
                  "comme pour l'opérateur.", size=15, color="lt2")
        for i, (v, l) in enumerate([("46 cm", "d'espace gantry–table"), ("60 cm", "sFOV"),
                                    ("85 cm", "FOV étendu¹")]):
            self.text(s, MX + i * 2.5, 5.95, 2.4, 0.4, v, size=20, bold=True, color="accent1")
            self.text(s, MX + i * 2.5, 6.35, 2.4, 0.3, l, size=10, color="lt2")
        self.variant_tag(s, "A · SIGNAL")
        self.chrome(s, show_section=False, sd={"footnotes": ["1. EFOV 85 cm réservé à la préparation et la planification de la "
                                      "radiothérapie ; non diagnostique."], "footnotes_y": 6.62,
                        "footnotes_w": 8.0})

    def a_versus(self):
        s = self.blank("A4")
        self.bg(s, "dk1")
        half = SW / 2
        self.rect(s, 0, 0, half, 6.72, fill="accent1")
        self.rect(s, half, 0, half, 6.72, fill="dk2")
        self.text(s, MX, 0.55, half - MX, 0.3, "PHILIPS", size=10, bold=True, color="dk1", spacing=300)
        self.text(s, MX, 0.85, half - MX - 0.4, 0.75, "Spectral CT 7500", size=34, bold=True, color="dk1")
        self.text(s, half + 0.6, 0.55, half - 1.2, 0.3, "GE HEALTHCARE", size=10, bold=True, color="accent4",
                  spacing=300, check=False)
        self.text(s, half + 0.6, 0.85, half - 1.2, 0.75, "Revolution Apex Elite", size=34, bold=True, color="lt1")
        vs = self.rect(s, half - 0.55, 0.85, 1.1, 1.1, fill="dk1", shape=MSO_SHAPE.OVAL)
        self.text(s, half - 0.55, 0.85, 1.1, 1.1, "VS", size=22, bold=True, color="lt1", align="c", anchor="m",
                  check=False)
        y = 2.3
        for lab, p, c, adv in VS_ROWS:
            self.text(s, MX, y, 3.0, 0.6, lab, size=12, color="dk1", anchor="m")
            self.text(s, MX + 2.9, y, half - MX - 3.65, 0.6, p, size=26, bold=True, color="dk1", align="r",
                      anchor="m", check=False)
            if adv:
                self.rect(s, half - 0.62, y + 0.15, 0.3, 0.3, fill="dk1", shape=MSO_SHAPE.OVAL)
                self.text(s, half - 0.62, y + 0.15, 0.3, 0.3, "✓", size=11, bold=True, color="accent1", align="c",
                          anchor="m", check=False)
            self.text(s, half + 0.6, y, 3.0, 0.6, c, size=26, bold=True, color="lt1", anchor="m", check=False)
            self.text(s, half + 3.2, y, 2.8, 0.6, lab, size=12, color="lt2", anchor="m", check=False)
            y += 0.68
        self.variant_tag(s, "A · SIGNAL")
        self.chrome(s, {}, show_section=False)

    # ============================================================ B « Suisse »
    def b_frame(self, s):
        self.bg(s, "lt1")
        self.grid(s, "lt2", alpha=0.45)

    def b_opener(self):
        s = self.blank("B1")
        self.b_frame(s)
        gw = (SW - 2 * MX) / 12
        self.rect(s, MX + 7 * gw, 0, SW - (MX + 7 * gw), SH, fill="dk1")
        self.image(s, "J-44", MX + 7 * gw, 0.9, SW - (MX + 7 * gw), 4.4, focus=(0.55, 0.5))
        self.text(s, MX, 0.62, 2, 0.4, "01", size=18, bold=True, color="dk1")
        self.text(s, MX + gw, 0.66, 4, 0.3, "SECTION", size=10, bold=True, color="accent6", spacing=300)
        self.set_title(s, "Gamme\nPhilips", y=1.75, w=6.4, h=3.0, size=88, color="dk1")
        self.text(s, MX, 5.0, 5.2, 0.9, "Des solutions innovantes, simples à utiliser et qui font une réelle "
                  "différence", size=16, color="dk2")
        for i, n in enumerate(["Verida", "Rembra", "CT 5300", "CT 5100", "CT 3500"]):
            self.text(s, MX + 7 * gw + 0.4, 5.55 + i * 0.0, 0.1, 0.1, "", check=False)
        self.text(s, MX + 7 * gw + 0.4, 5.6, 5, 0.9, [[("Chapitres", {"bold": True, "color": "accent1"})],
                                                     ["Verida · Rembra · CT 5300 · CT 5100 · CT 3500"]],
                  size=13, color="lt1")
        self.variant_tag(s, "B · SUISSE", dark=False)
        self.chrome(s, {}, dark=False, show_section=False)

    def b_bento(self):
        s = self.blank("B2")
        self.b_frame(s)
        self.text(s, MX, 0.62, 4, 0.3, "GAMME PHILIPS", size=10, bold=True, color="accent6", spacing=300)
        self.set_title(s, "Trois familles, une même intelligence", y=0.95, size=34, h=0.7, color="dk1")
        gw = (SW - 2 * MX) / 12
        cols = [(0, 4), (4, 5), (9, 3)]
        order = [GAMME[0], GAMME[1], GAMME[2]]
        for (c0, span), (fam, pillar, items) in zip(cols, order):
            x, w = MX + c0 * gw + 0.08, span * gw - 0.16
            head = "dk1" if fam != "Radiothérapie" else "accent6"
            self.rect(s, x, 1.95, w, 0.55, fill=head)
            self.text(s, x + 0.2, 1.95, w - 0.4, 0.55, fam, size=14, bold=True, color="lt1", anchor="m")
            self.text(s, x, 2.6, w, 0.4, pillar, size=10, italic=True, color="dk2")
            y = 3.1
            for name, new, img in items:
                ln = self.rect(s, x, y, w, 0.008, fill="dk1")
                if img and fam != "Radiothérapie":
                    self.image(s, img, x, y + 0.08, 1.15, 0.62, focus=(0.5, 0.55))
                    tx = x + 1.3
                else:
                    tx = x
                self.text(s, tx, y + 0.12, w - (tx - x), 0.32, name, size=14, bold=True, color="dk1")
                if new:
                    self.text(s, tx, y + 0.44, 1.2, 0.25, "NOUVEAU", size=9, bold=True, color="accent6",
                              spacing=150, check=False)
                y += 0.82
        self.variant_tag(s, "B · SUISSE", dark=False)
        self.chrome(s, {}, dark=False, show_section=False)

    def b_stat(self):
        s = self.blank("B3")
        self.b_frame(s)
        gw = (SW - 2 * MX) / 12
        self.text(s, MX, 0.62, 5, 0.3, "REMBRA · NOUVEAU", size=10, bold=True, color="accent6", spacing=300)
        self.text(s, MX - 0.08, 1.0, 7.5, 3.4, [[("85", {"size": 230, "bold": True, "color": "dk1"}),
                                                  (" cm", {"size": 60, "bold": True, "color": "accent6"})]],
                  check=False)
        self.set_title(s, "Le tunnel le plus large de sa catégorie", y=4.5, w=6.6, h=0.6, size=24, color="dk1")
        self.image(s, "J-50", MX + 8 * gw, 0.95, 4 * gw, 3.6, focus=(0.5, 0.5))
        y = 5.3
        for i, (v, l) in enumerate([("46 cm", "d'espace gantry–table"), ("60 cm", "sFOV"),
                                    ("85 cm", "FOV étendu¹"), ("0,34 s", "temps de rotation")]):
            x = MX + i * 3 * gw
            self.rect(s, x, y, 3 * gw - 0.15, 0.01, fill="dk1")
            self.text(s, x, y + 0.12, 3 * gw - 0.2, 0.45, v, size=22, bold=True, color="dk1")
            self.text(s, x, y + 0.6, 3 * gw - 0.2, 0.3, l, size=10, color="dk2")
        self.variant_tag(s, "B · SUISSE", dark=False)
        self.chrome(s, show_section=False, dark=False, sd={"footnotes": ["1. EFOV 85 cm réservé à la préparation et la planification de la "
                                      "radiothérapie ; non diagnostique."], "footnotes_y": 6.55})

    def b_versus(self):
        s = self.blank("B4")
        self.b_frame(s)
        self.text(s, MX, 0.62, 5, 0.3, "COMPÉTITION · GE HEALTHCARE", size=10, bold=True, color="accent6",
                  spacing=300)
        self.set_title(s, "Spectral CT 7500 vs Revolution Apex Elite", y=0.95, size=30, h=0.7, color="dk1")
        x0, wl, wc = MX, 4.0, 3.9
        xp, xc = x0 + wl, x0 + wl + wc
        self.rect(s, xp, 1.95, wc, 4.65, fill="accent3")
        self.text(s, xp + 0.25, 2.0, wc - 0.5, 0.6, [[("Spectral CT 7500", {"bold": True, "size": 16})],
                                                       [("Philips · Premium", {"size": 10})]], color="dk1",
                  anchor="m", check=False)
        self.text(s, xc + 0.25, 2.0, wc - 0.5, 0.6, [[("Revolution Apex Elite", {"bold": True, "size": 16})],
                                                       [("GE HealthCare", {"size": 10})]], color="dk2",
                  anchor="m", check=False)
        y = 2.75
        for lab, p, c, adv in VS_ROWS:
            self.rect(s, x0, y, wl + 2 * wc, 0.008, fill="dk1")
            self.text(s, x0, y, wl, 0.6, lab, size=13, color="dk1", anchor="m")
            self.text(s, xp + 0.25, y, wc - 1, 0.6, p, size=20, bold=True, color="dk1", anchor="m", check=False)
            if adv:
                self.rect(s, xp + wc - 0.6, y + 0.15, 0.3, 0.3, fill="dk1", shape=MSO_SHAPE.OVAL)
                self.text(s, xp + wc - 0.6, y + 0.15, 0.3, 0.3, "✓", size=11, bold=True, color="lt1",
                          align="c", anchor="m", check=False)
            self.text(s, xc + 0.25, y, wc - 0.5, 0.6, c, size=20, color="dk2", anchor="m", check=False)
            y += 0.62
        self.variant_tag(s, "B · SUISSE", dark=False)
        self.chrome(s, {}, dark=False, show_section=False)

    # ============================================================ C « Anneau »
    def c_frame(self, s, cx=9.6, cy=3.6):
        self.bg(s, "dk1")
        glow = self.rect(s, cx - 5.5, cy - 5.5, 11, 11, fill="dk2", shape=MSO_SHAPE.OVAL)
        sp = glow._element.spPr
        for el in sp.findall(qn("a:solidFill")):
            sp.remove(el)
        g = etree.SubElement(sp, qn("a:gradFill"), rotWithShape="1")
        lst = etree.SubElement(g, qn("a:gsLst"))
        for pos, col, a in [(0, "accent1/dk1/0.55", 0.55), (60000, "dk2", 0.35), (100000, "dk1", 0)]:
            gs = etree.SubElement(lst, qn("a:gs"), pos=str(pos))
            c = etree.SubElement(gs, qn("a:srgbClr"), val=self.P[col])
            etree.SubElement(c, qn("a:alpha"), val=str(int(a * 100000)))
        path = etree.SubElement(g, qn("a:path"), path="circle")
        etree.SubElement(path, qn("a:fillToRect"), l="50000", t="50000", r="50000", b="50000")
        ln = sp.find(qn("a:ln"))
        if ln is not None:
            sp.remove(g)
            ln.addprevious(g)

    def c_opener(self):
        s = self.blank("C1")
        cx, cy = 9.3, 3.55
        self.c_frame(s, cx, cy)
        self.rings(s, cx, cy, [2.45, 2.85, 3.3, 3.85, 4.5], alpha=0.22)
        pic = self.image(s, "J-44", cx - 2.2, cy - 2.2, 4.4, 4.4, focus=(0.53, 0.5))
        self.oval_pic(pic)
        r = self.rect(s, cx - 2.25, cy - 2.25, 4.5, 4.5, line="accent1", shape=MSO_SHAPE.OVAL, line_w=2.5)
        self.glow(r, rad=14, alpha=0.6)
        t = self.text(s, MX, 1.0, 3, 1.2, "01", size=72, bold=True, color="accent1", check=False)
        self.text_glow(t)
        self.set_title(s, "Gamme\nPhilips", y=2.3, w=5.6, h=2.4, size=66)
        self.text(s, MX, 4.55, 5.0, 0.9, "Des solutions innovantes, simples à utiliser et qui font une réelle "
                  "différence", size=15, color="lt2")
        cx2 = MX
        for n in ["Verida", "Rembra", "CT 5300", "CT 5100", "CT 3500"]:
            w = text_width(n, 11, True) + 0.45
            c = self.rect(s, cx2, 5.65, w, 0.36, line="accent1", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            self.alpha_line(c, 0.7)
            self.text(s, cx2, 5.65, w, 0.36, n, size=11, bold=True, align="c", anchor="m", check=False)
            cx2 += w + 0.12
        self.variant_tag(s, "C · ANNEAU")
        self.chrome(s, {}, show_section=False)

    def c_bento(self):
        s = self.blank("C2")
        self.c_frame(s, SW / 2, 4.2)
        self.rings(s, SW / 2, 4.2, [3.0, 4.2, 5.6], alpha=0.12)
        self.label(s, MX, 0.55, "Gamme Philips")
        self.set_title(s, "Trois familles, une même intelligence", y=0.9, size=32, h=0.7)
        cw = (SW - 2 * MX - 0.5) / 3
        for i, (fam, pillar, items) in enumerate(GAMME):
            x = MX + i * (cw + 0.25)
            card = self.rect(s, x, 1.95, cw, 4.6, fill="dk2", line="accent1", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                             radius=0.05, line_w=1)
            self.alpha_line(card, 0.45)
            self.alpha_fill(card, 0.7)
            hero = items[0][2] or None
            if hero:
                pic = self.image(s, hero, x + cw / 2 - 0.85, 2.15, 1.7, 1.7, focus=(0.5, 0.45))
                self.oval_pic(pic)
                ring = self.rect(s, x + cw / 2 - 0.9, 2.1, 1.8, 1.8, line="accent1", shape=MSO_SHAPE.OVAL,
                                 line_w=1.5)
                self.glow(ring, rad=10, alpha=0.5)
            else:
                self.rings(s, x + cw / 2, 3.0, [0.55, 0.75, 0.9], alpha=0.5)
                self.text(s, x + cw / 2 - 0.85, 2.75, 1.7, 0.5, "RT", size=20, bold=True, color="accent1",
                          align="c", anchor="m", check=False)
            self.text(s, x + 0.25, 4.0, cw - 0.5, 0.4, fam, size=17, bold=True, align="c")
            self.text(s, x + 0.25, 4.4, cw - 0.5, 0.35, pillar, size=10, color="lt2", italic=True, align="c")
            y = 4.9
            for name, new, _ in items:
                self.text(s, x + 0.35, y, cw - 1.7, 0.3, name, size=12, bold=True)
                if new:
                    b = self.badge_new(s, x + cw - 1.3, y)
                    self.glow(b, rad=8, alpha=0.5)
                y += 0.38
        self.variant_tag(s, "C · ANNEAU")
        self.chrome(s, {}, show_section=False)

    def c_stat(self):
        s = self.blank("C3")
        cx, cy = 4.3, 3.7
        self.c_frame(s, cx, cy)
        self.rings(s, cx, cy, [2.2, 2.6, 3.05, 3.6], alpha=0.25)
        r = self.rect(s, cx - 1.85, cy - 1.85, 3.7, 3.7, line="accent1", shape=MSO_SHAPE.OVAL, line_w=3)
        self.glow(r, rad=20, alpha=0.55)
        t = self.text(s, cx - 1.85, cy - 0.95, 3.7, 1.5, "85", size=110, bold=True, color="lt1", align="c",
                      anchor="m", check=False)
        self.text_glow(t, rad=12, alpha=0.5)
        self.text(s, cx - 1.85, cy + 0.5, 3.7, 0.5, "cm d'ouverture", size=18, bold=True, color="accent1",
                  align="c")
        self.label(s, 8.0, 1.2, "Rembra · Nouveau")
        self.set_title(s, "Le tunnel le plus large de sa catégorie", x=8.0, y=1.6, w=4.8, h=1.3, size=28)
        self.text(s, 8.0, 3.0, 4.6, 0.9, "Confort, accessibilité et flexibilité d'utilisation pour le patient "
                  "comme pour l'opérateur.", size=14, color="lt2")
        for i, (v, l) in enumerate([("46 cm", "d'espace gantry–table"), ("60 cm", "sFOV"),
                                    ("85 cm", "FOV étendu¹")]):
            y = 4.15 + i * 0.68
            dot = self.rect(s, 8.0, y + 0.12, 0.22, 0.22, fill="accent1", shape=MSO_SHAPE.OVAL)
            self.glow(dot, rad=6, alpha=0.6)
            self.text(s, 8.4, y, 1.6, 0.45, v, size=20, bold=True, color="accent1")
            self.text(s, 10.0, y + 0.05, 2.8, 0.4, l, size=12, color="lt2", anchor="m")
        self.label(s, MX, 0.55, "Gamme Philips")
        self.variant_tag(s, "C · ANNEAU")
        self.chrome(s, show_section=False, sd={"footnotes": ["1. EFOV 85 cm réservé à la préparation et la planification de la "
                                      "radiothérapie ; non diagnostique."], "footnotes_y": 6.55})

    def c_versus(self):
        s = self.blank("C4")
        cx = SW / 2
        self.c_frame(s, cx, 1.35)
        self.rings(s, cx, 1.35, [0.75, 0.95, 1.2], alpha=0.35)
        vs = self.rect(s, cx - 0.6, 0.75, 1.2, 1.2, fill="dk1", line="accent1", shape=MSO_SHAPE.OVAL, line_w=2)
        self.glow(vs, rad=14, alpha=0.6)
        self.text(s, cx - 0.6, 0.75, 1.2, 1.2, "VS", size=24, bold=True, align="c", anchor="m", check=False)
        t = self.text(s, MX, 0.85, cx - MX - 1.0, 0.9, [[("Spectral CT 7500", {"bold": True, "size": 28})],
                                                         [("Philips · Premium", {"size": 11, "color": "lt2"})]],
                      color="accent1", align="r", check=False)
        self.text_glow(t, rad=8, alpha=0.35)
        self.text(s, cx + 1.0, 0.85, cx - MX - 1.0, 0.9, [[("Revolution Apex Elite", {"bold": True, "size": 28})],
                                                           [("GE HealthCare", {"size": 11, "color": "lt2"})]],
                  color="accent4", check=False)
        y = 2.35
        for lab, p, c, adv in VS_ROWS:
            row = self.rect(s, MX, y, SW - 2 * MX, 0.58, fill="dk2", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            self.alpha_fill(row, 0.55)
            self.text(s, cx - 1.6, y, 3.2, 0.58, lab, size=12, color="lt2", align="c", anchor="m", check=False)
            pt = self.text(s, MX + 0.4, y, cx - MX - 2.2, 0.58, p, size=22, bold=True, color="accent1", align="r",
                           anchor="m", check=False)
            if adv:
                self.text_glow(pt, rad=8, alpha=0.45)
                d = self.rect(s, MX + 0.25, y + 0.15, 0.28, 0.28, fill="accent1", shape=MSO_SHAPE.OVAL)
                self.text(s, MX + 0.25, y + 0.15, 0.28, 0.28, "✓", size=10, bold=True, color="dk1", align="c",
                          anchor="m", check=False)
            self.text(s, cx + 1.8, y, cx - MX - 2.2, 0.58, c, size=22, bold=True, color="lt1", anchor="m",
                      check=False)
            y += 0.7
        self.variant_tag(s, "C · ANNEAU")
        self.chrome(s, {}, show_section=False)

    def run(self):
        # retirer la couverture d'origine : les prévisualisations ne contiennent que les variantes
        sldIdLst = self.prs.slides._sldIdLst
        first = sldIdLst[0]
        self.prs.part.drop_rel(first.rId)
        sldIdLst.remove(first)
        for v in "abc":
            for kind in ("opener", "bento", "stat", "versus"):
                getattr(self, f"{v}_{kind}")()
        self.prs.save(self.out_path)
        print(f"{len(self.prs.slides)} slides → {self.out_path}")
        for w in self.warnings:
            print("ATTENTION", w)


if __name__ == "__main__":
    Previews().run()
